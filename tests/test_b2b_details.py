# -*- coding: utf-8 -*-
"""Офлайн-тесты карточки и документации B2B — реальный браузер не запускается."""
import json
import os
import zipfile
from pathlib import Path

import openpyxl
import pytest
import requests

from tenders import b2b_details as d
from tenders import core


def quiet(*_):
    pass


def aggregate(number="101"):
    return {
        "procedure": {
            "number": number, "title": "Муфты", "url": "",
            "type": "Запрос предложений", "status": "Объявлена",
            "published_at": "2026-10-06T10:00:00+03:00",
            "modified_at": "", "deadline": "2026-10-09T10:00:00+03:00",
            "positions_count": 2, "currency": "RUB", "total_price": "",
            "category": "",
        },
        "organizer": {
            "name": "ООО Организатор", "full_name": "ООО Организатор",
            "inn": "123", "kpp": "456", "ogrn": "789", "okpo": "012",
            "legal_address": "Москва", "profile_url": "https://example/firms/1",
        },
        "terms": {
            "held_for": "", "payment": "Постоплата", "delivery": "Доставка",
            "delivery_address": ["Склад"], "submission_method": "",
            "participant_comment": "",
        },
        "extra_fields": {},
    }


def payload(number=101):
    return {
        "status": {"code": 1, "message": "success"},
        "trade_aggregate": {
            "positions_count": 2,
            "trade": {
                "id": number,
                "type": {"name": {"hint": {"title": "Запрос предложений"}}},
                "status": {"name": {"hint": {"title": "Объявлена"}}},
                "date_published": "2026-10-06T10:00:00+03:00",
                "date_modified": "2026-10-06T11:00:00+03:00",
                "firm": {
                    "short_name": "ООО Организатор", "full_name": "Общество",
                    "inn": "123", "kpp": "456", "ogrn": "789", "okpo": "012",
                    "jury_address": "Москва", "url": "https://example/firms/1",
                },
                "fields_values": {
                    "subject": {"value": "Муфты"},
                    "payment_terms": {"value": "Постоплата"},
                    "delivery_terms": {"value": "До склада"},
                    "delivery_address": [{"address": {"address_string": "Склад 1"}}],
                    "currency": {"value": "RUB"},
                },
                "settings_values": {
                    "main_stage_date_end": {"value": "2026-10-09T10:00:00+03:00"},
                },
            },
        },
    }


def browser_data(layout="market-next"):
    return {
        "layout": layout,
        "title": "Муфты по чертежу",
        "number": "№ 101",
        "url": "https://www.b2b-center.ru/app/market-next/x/tender-101/",
        "category": "Изделия по чертежам",
        "status": "Приём заявок",
        "sections": {
            "Организатор": "ООО Организатор",
            "Процедура проводится": "Для собственных нужд",
            "Условия оплаты": "100% постоплата",
            "Условия поставки / оказания услуг": "До склада",
            "Адрес поставки / оказания услуг": "Склад 1",
            "Способ проведения": "По всем позициям",
            "Цены подаются": "Без НДС",
            "Аналоги": "Разрешены",
            "Прочее поле": "Значение",
        },
        "contacts": "Контакты Иванов Иван, специалист, ivanov@example.ru, +7 900 111-22-33",
        "positions": [
            {"lot": "Лот 1", "number": "1.1", "name": "Муфта",
             "raw": {"Количество": "2", "Единица измерения": "Штука"}},
            {"lot": "Лот 2", "number": "2.1", "name": "Втулка",
             "raw": {"Количество": "1", "Единица измерения": "Штука"}},
        ],
        "documents": {
            "package": None,
            "listed": [{"name": "ТЗ.pdf", "url": "https://example/download",
                        "checksum": "abc", "size": "10 Кб"}],
            "obsolete": [], "fingerprint": "fp",
        },
        "errors": [],
        "files_extracted": 0,
        "archives_downloaded": 0,
        "skipped_unchanged": 0,
    }


@pytest.fixture
def workbook(tmp_path, monkeypatch):
    work = tmp_path / "work"
    work.mkdir()
    monkeypatch.setattr(core, "WORK_DIR", str(work))
    path = tmp_path / "заметки.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = core.SHEET
    ws.append(core.HEADERS)
    for row, (number, etp, link) in enumerate([
        ("101", "b2b-center", "https://www.b2b-center.ru/app/market-next/x/tender-101/"),
        ("102", "Фабрикант", "https://example/102"),
        ("103", "b2b-center", "https://www.b2b-center.ru/market/x/tender-103/"),
    ], start=2):
        ws.cell(row, core.COL_CUSTOMER, "Заказчик")
        ws.cell(row, core.COL_NUMBER, number)
        ws.cell(row, core.COL_ETP, etp)
        ws.cell(row, core.COL_LINK, "Открыть")
        ws.cell(row, core.COL_LINK).hyperlink = link
    wb.save(path)
    return {"xlsx": str(path), "work": work}


def test_normalize_aggregate():
    result = d.normalize_aggregate(payload(), "101", "https://example/tender-101/")

    assert result["procedure"]["title"] == "Муфты"
    assert result["procedure"]["deadline"].startswith("2026-10-09")
    assert result["organizer"]["inn"] == "123"
    assert result["organizer"]["legal_address"] == "Москва"
    assert result["terms"]["delivery_address"] == ["Склад 1"]


def test_normalize_aggregate_rejects_other_number():
    with pytest.raises(d.ProtocolChanged):
        d.normalize_aggregate(payload(999), "101", "https://example")


@pytest.mark.parametrize("layout", ["market-next", "classic"])
def test_build_record_for_both_layouts(layout):
    result = d.build_record("101", "https://example", aggregate(), browser_data(layout))

    assert result["meta"]["layout"] == layout
    assert result["meta"]["complete"] is True
    assert result["procedure"]["title"] == "Муфты по чертежу"
    assert result["terms"]["held_for"] == "Для собственных нужд"
    assert result["requirements"]["Аналоги"] == "Разрешены"
    assert result["extra_fields"]["Прочее поле"] == "Значение"
    assert len(result["positions"]) == 2
    assert result["contacts"][0]["emails"] == ["ivanov@example.ru"]


def test_build_record_marks_partial_result():
    data = browser_data()
    data["errors"] = ["Позиции не загрузились"]
    assert d.build_record("101", "https://example", aggregate(), data)["meta"]["complete"] is False


def test_safe_extract_zip(tmp_path):
    archive = tmp_path / "docs.zip"
    with zipfile.ZipFile(archive, "w") as zf:
        zf.writestr("ТЗ/чертёж.txt", "данные")

    extracted = d.safe_extract_zip(archive, tmp_path / "out")

    assert extracted == [os.path.join("ТЗ", "чертёж.txt")]
    assert (tmp_path / "out" / "ТЗ" / "чертёж.txt").read_text(encoding="utf-8") == "данные"


@pytest.mark.parametrize("unsafe", ["../outside.txt", "/absolute.txt", "C:/windows.txt"])
def test_safe_extract_zip_rejects_escape(tmp_path, unsafe):
    archive = tmp_path / "bad.zip"
    with zipfile.ZipFile(archive, "w") as zf:
        zf.writestr("good.txt", "good")
        zf.writestr(unsafe, "bad")

    with pytest.raises(d.DetailsError):
        d.safe_extract_zip(archive, tmp_path / "out")
    assert not (tmp_path / "out" / "good.txt").exists()
    assert not (tmp_path / "outside.txt").exists()


def test_safe_extract_zip_rejects_corrupted_archive(tmp_path):
    archive = tmp_path / "bad.zip"
    archive.write_bytes(b"not a zip")

    with pytest.raises(d.DetailsError):
        d.safe_extract_zip(archive, tmp_path / "out")


def test_atomic_json_preserves_old_file_on_error(tmp_path):
    path = tmp_path / d.INFO_FILENAME
    path.write_text("старое", encoding="utf-8")

    with pytest.raises(TypeError):
        d.atomic_write_json(path, {"bad": object()})

    assert path.read_text(encoding="utf-8") == "старое"
    assert not (tmp_path / (d.INFO_FILENAME + ".tmp")).exists()


def test_document_fingerprint_is_order_independent():
    docs = [
        {"name": "a", "url": "u1", "checksum": "1"},
        {"name": "b", "url": "u2", "checksum": "2"},
    ]
    assert d.document_fingerprint(docs) == d.document_fingerprint(list(reversed(docs)))


def test_login_can_finish_in_new_browser_tab():
    class Locator:
        def __init__(self, count):
            self._count = count

        def count(self):
            return self._count

    class Page:
        def __init__(self, logged_in, url):
            self.logged_in = logged_in
            self.url = url
            self.visited = []

        def is_closed(self):
            return False

        def locator(self, _selector):
            return Locator(1 if self.logged_in else 0)

        def goto(self, url, **_kwargs):
            self.visited.append(url)
            self.url = url

    original = Page(False, "https://www.b2b-center.ru/login")
    authenticated = Page(True, "https://www.b2b-center.ru/app/next/main/")
    browser = d.PlaywrightBrowser(quiet)
    browser.page = original
    browser.context = type("Context", (), {"pages": [original, authenticated]})()

    browser._ensure_login("https://www.b2b-center.ru/app/market-next/x/tender-101/")

    assert browser.page is authenticated
    assert authenticated.visited == [
        "https://www.b2b-center.ru/app/market-next/x/tender-101/"]


class FakeBrowser:
    def __init__(self, say, data=None, error=None, make_package=False):
        self.data = data or browser_data()
        self.error = error
        self.make_package = make_package
        self.calls = []

    def __enter__(self):
        return self

    def __exit__(self, *_):
        return None

    def collect(self, number, url, tender_dir, previous=None):
        self.calls.append((number, url, tender_dir, previous))
        if self.error:
            raise self.error
        data = json.loads(json.dumps(self.data, ensure_ascii=False))
        data["number"] = number
        data["url"] = url
        if self.make_package:
            etp = Path(tender_dir) / core.SUBFOLDER
            archive = etp / "Документы.zip"
            with zipfile.ZipFile(archive, "w") as zf:
                zf.writestr("ТЗ.txt", "техническое задание")
            extracted = d.safe_extract_zip(archive, etp)
            data["documents"]["package"] = {
                "name": archive.name, "path": os.path.join(core.SUBFOLDER, archive.name),
                "type": "zip", "status": "downloaded", "extracted": extracted,
            }
            data["archives_downloaded"] = 1
            data["files_extracted"] = len(extracted)
        return data


def fake_factory(fake):
    return lambda _say: fake


def test_collect_exact_number_writes_txt_and_files(workbook, monkeypatch):
    monkeypatch.setattr(d, "fetch_trade_aggregate", lambda _session, number: aggregate(number))
    fake = FakeBrowser(quiet, make_package=True)

    result = d.collect_b2b_details(
        workbook["xlsx"], ["101"], say=quiet, browser_factory=fake_factory(fake))

    tender = workbook["work"] / "101"
    saved = json.loads((tender / d.INFO_FILENAME).read_text(encoding="utf-8"))
    assert result["processed"] == ["101"]
    assert result["txt_written"] == 1
    assert result["archives_downloaded"] == 1
    assert result["files_extracted"] == 1
    assert saved["procedure"]["number"] == "101"
    assert saved["positions"][1]["name"] == "Втулка"
    assert (tender / core.SUBFOLDER / "Документы.zip").exists()
    assert (tender / core.SUBFOLDER / "ТЗ.txt").exists()
    assert [call[0] for call in fake.calls] == ["101"]


def test_collect_rejects_other_etp_and_unknown_number(workbook):
    result = d.collect_b2b_details(workbook["xlsx"], ["102", "999"], say=quiet)

    assert result["processed"] == []
    assert result["errors"] == [
        {"number": "102", "error": "у строки другая ЭТП"},
        {"number": "999", "error": "номер не найден в Excel"},
    ]


def test_collect_dry_run_does_not_open_browser_or_create_folder(workbook):
    def forbidden(_say):
        raise AssertionError("браузер не должен запускаться")

    result = d.collect_b2b_details(
        workbook["xlsx"], ["101"], dry_run=True, say=quiet,
        browser_factory=forbidden)

    assert result["dry_run"] is True
    assert not (workbook["work"] / "101").exists()


def test_collect_auth_timeout_preserves_old_txt(workbook, monkeypatch):
    monkeypatch.setattr(d, "fetch_trade_aggregate", lambda _session, number: aggregate(number))
    tender = workbook["work"] / "101"
    (tender / core.SUBFOLDER).mkdir(parents=True)
    info = tender / d.INFO_FILENAME
    info.write_text("старое", encoding="utf-8")
    fake = FakeBrowser(quiet, error=d.AuthRequired("нужен вход"))

    result = d.collect_b2b_details(
        workbook["xlsx"], ["101"], say=quiet, browser_factory=fake_factory(fake))

    assert result["auth_required"] is True
    assert result["stopped"] == "нужен вход"
    assert info.read_text(encoding="utf-8") == "старое"


def test_collect_protocol_change_stops_and_preserves_old_txt(workbook, monkeypatch):
    monkeypatch.setattr(d, "fetch_trade_aggregate", lambda _session, number: aggregate(number))
    tender = workbook["work"] / "101"
    (tender / core.SUBFOLDER).mkdir(parents=True)
    info = tender / d.INFO_FILENAME
    info.write_text("старое", encoding="utf-8")
    fake = FakeBrowser(quiet, error=d.ProtocolChanged("структура изменилась"))

    result = d.collect_b2b_details(
        workbook["xlsx"], ["101"], say=quiet, browser_factory=fake_factory(fake))

    assert result["stopped"] == "структура изменилась"
    assert info.read_text(encoding="utf-8") == "старое"


def test_collect_network_error_keeps_browser_data_as_partial(workbook, monkeypatch):
    def offline(_session, _number):
        raise requests.ConnectionError("offline")

    monkeypatch.setattr(d, "fetch_trade_aggregate", offline)
    fake = FakeBrowser(quiet)

    result = d.collect_b2b_details(
        workbook["xlsx"], ["101"], say=quiet, browser_factory=fake_factory(fake))

    saved = json.loads((workbook["work"] / "101" / d.INFO_FILENAME)
                       .read_text(encoding="utf-8"))
    assert result["processed"] == ["101"]
    assert saved["meta"]["complete"] is False
    assert "JSON карточки" in saved["meta"]["errors"][0]


def test_collect_reports_unchanged_files(workbook, monkeypatch):
    monkeypatch.setattr(d, "fetch_trade_aggregate", lambda _session, number: aggregate(number))
    data = browser_data()
    data["skipped_unchanged"] = 3
    fake = FakeBrowser(quiet, data=data)

    result = d.collect_b2b_details(
        workbook["xlsx"], ["101"], say=quiet, browser_factory=fake_factory(fake))

    assert result["skipped_unchanged"] == 3
