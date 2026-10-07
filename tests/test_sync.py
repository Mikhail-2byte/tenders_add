# -*- coding: utf-8 -*-
"""Тесты ядра sync(): добавление тендеров, сортировка и сохранность данных.

Это самая рискованная часть программы — она стирает всю область данных листа и
пишет её заново в новом порядке, поэтому проверяем не только «добавилось», но и
«ничего не пропало» и «обвязка листа осталась рабочей».

Запуск: py -3.12 -m pytest
"""
import datetime
import os
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
import pytest
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.table import Table

from tenders import b2b_links
from tenders import b2b_details
from tenders import core as at


# Строки листа «Справочники»: статус, порядок сортировки, номер группы.
REF_ROWS = [("Интересный", "В работе", 1), ("В работе", "Податься", 1),
            ("Податься", "Переторжка", 1), ("Подался", "Интересный", 2),
            ("Переторжка", "Подался", 3), ("Ожидаем итоги", "Ожидаем итоги", 4),
            ("Победа", "Победа", 5), ("Проиграли", "Проиграли", 6),
            ("Не участвуем", "Не участвуем", 7), ("Отменён", "Отменён", 8)]

# Строки данных: номер, заказчик, статус, срок, пометка в «Действие».
DATA_ROWS = [
    ("111", "Первый заказчик", "Подался", datetime.datetime(2026, 7, 1, 9, 0),
     "Прежний статус: Податься"),
    ("222", "Второй заказчик", "В работе", datetime.datetime(2026, 9, 20, 16, 0), None),
    ("333", "Третий заказчик", "Проиграли", None, None),
    ("444", "Четвёртый заказчик", "В работе", datetime.datetime(2026, 7, 5, 12, 0), None),
    ("555\n556", "Пятый заказчик", "Не участвуем", None, None),
]

# Выпадающий список статусов, как его пишет Excel: расширение x14 внутри <extLst>.
# Атрибут xr:uid оставлен намеренно — read_ext_lst обязан его выкусить.
EXT_SAMPLE = (
    '<extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}"'
    ' xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"'
    ' xmlns:xr="http://schemas.microsoft.com/office/spreadsheetml/2014/revision">'
    '<x14:dataValidations count="1"'
    ' xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
    '<x14:dataValidation type="list" allowBlank="1" showInputMessage="1"'
    ' showErrorMessage="1" errorTitle="Недопустимый статус"'
    ' error="Выберите статус из списка." xr:uid="{9DFE99AE-C652-4584-A01C-B6A549E1879D}">'
    '<x14:formula1><xm:f>Справочники!$A$2:$A$11</xm:f></x14:formula1>'
    '<xm:sqref>I2:I1994</xm:sqref>'
    '</x14:dataValidation></x14:dataValidations></ext></extLst>'
).encode()


def line_1c(number, customer="ООО Ромашка", deadline="10.08.2026 12:30",
            status="В работе", name="наименование", etp="Сбербанк-АСТ",
            comment="комментарий из 1С"):
    """Строка выгрузки 1С: 13 полей через TAB, номер — поле 5, заказчик — поле 6."""
    return "\t".join(["1", "01.08.2026", deadline, "10", number, customer,
                      "Тендер 000000000025397 от 14.09.2026 11:54:11", "", "Нет",
                      status, name, etp, comment])


@pytest.fixture
def env(tmp_path, monkeypatch):
    """Временный xlsx в новом макете + временная папка работ; боевое не трогаем."""
    work = tmp_path / "В работе"
    work.mkdir()
    monkeypatch.setattr(at, "WORK_DIR", str(work))
    monkeypatch.setattr(at, "HISTORY_PATH", str(tmp_path / "history.jsonl"))
    monkeypatch.setattr(at, "LOG_PATH", str(tmp_path / "add_tenders.log"))

    xlsx = tmp_path / "заметки.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = at.SHEET
    ws.append(at.HEADERS)

    for i, (number, customer, status, deadline, action) in enumerate(DATA_ROWS):
        r = i + 2
        ws.cell(r, at.COL_CUSTOMER, customer)
        ws.cell(r, at.COL_NAME, "наименование " + number.split("\n")[0])
        cell = ws.cell(r, at.COL_NUMBER, number)
        cell.number_format = "@"
        ws.cell(r, at.COL_STATUS, status)
        if action:
            ws.cell(r, at.COL_ACTION, action)
        cell = ws.cell(r, at.COL_DEADLINE, deadline)
        cell.number_format = at.DEADLINE_FORMAT      # формат ставится ПОСЛЕ значения
        ws.cell(r, at.COL_SORT, at.SORT_FORMULA.format(row=r))
        ws.cell(r, at.COL_SORT).font = Font(name="Calibri", sz=8, color="FFBFBFBF")
        ws.cell(r, at.COL_NUMBER).font = Font(name="Calibri", sz=10)
        ws.row_dimensions[r].height = at.ROW_HEIGHT

    # Ручные данные и оформление, которые обязаны пережить перезапись листа.
    ws["D2"].value = "Открыть"
    ws["D2"].hyperlink = "https://example.com/1"
    ws["D3"].value = "Открыть"
    # «Сломанная» ссылка: адрес лежит в location, а не в target — таких в файле много.
    ws["D3"].hyperlink = Hyperlink(ref="D3", location="https://example.com/broken")
    ws["D3"].font = Font(bold=True, color="FFFF0000")
    ws["G2"].value = "ручной комментарий"
    ws["G2"].comment = Comment("важное примечание", "менеджер")
    ws["J2"].value = "Иванов, 8-900"
    # Остаточная «зебра»: строка-шаблон должна браться из НЕзалитой строки.
    for c in range(1, at.NCOL):
        ws.cell(2, c).fill = PatternFill("solid", fgColor="FFE8EEF7")

    last = len(DATA_ROWS) + 1
    ws.add_table(Table(displayName=at.TABLE, ref="A1:L%d" % last))
    ws.conditional_formatting.add("F2:F%d" % last, FormulaRule(
        formula=['AND(ISNUMBER($F2),$F2<TODAY())'],
        fill=PatternFill("solid", bgColor="FFFF7C80")))
    ws.conditional_formatting.add("A2:K%d" % last, FormulaRule(
        formula=['$I2="Проиграли"'], fill=PatternFill("solid", bgColor="FFFFC7CE")))

    ref = wb.create_sheet(at.REF_SHEET)
    ref.append(["Статус", None, "Порядок сортировки", None, "Группа"])
    for status, order, group in REF_ROWS:
        ref.append([status, None, order, None, group])

    dash = wb.create_sheet("Дашборд")
    dash["B2"] = '=COUNTIF(Тендеры!$I$2:$I$%d,"Победа")' % last

    wb.save(xlsx)
    # Дошиваем выпадающий список — openpyxl его писать не умеет.
    at.write_ext_lst(str(xlsx), 1, EXT_SAMPLE, last)
    return {"xlsx": str(xlsx), "work": str(work), "tmp": tmp_path}


def load(env):
    return openpyxl.load_workbook(env["xlsx"])[at.SHEET]


def numbers(ws):
    """Номера тендеров сверху вниз (столбец C)."""
    out = []
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, at.COL_NUMBER).value
        if v not in (None, ""):
            out.append(str(v))
    return out


def row_of(ws, number):
    """Строка листа, где лежит номер (позиция плавает — ищем по значению)."""
    return 2 + numbers(ws).index(number)


def sheet_xml(env):
    with zipfile.ZipFile(env["xlsx"]) as z:
        return z.read("xl/worksheets/sheet1.xml")


def quiet(*_):
    """Заглушка вместо print — тесты не должны шуметь в выводе."""


# ------------------------- порядок строк -------------------------

def test_rows_sorted_by_group_then_deadline(env):
    """Лист пересортирован: группа статуса, внутри группы — срок подачи."""
    at.sync(env["xlsx"], line_1c("444"), say=quiet)

    # 444 и 222 — группа 1 (по сроку), 111 — группа 3, 333 — 6, 555 — 7.
    assert numbers(load(env)) == ["444", "222", "111", "333", "555\n556"]


def test_new_row_placed_by_sort_not_appended(env):
    """Новый тендер с ближайшим сроком встаёт строкой 2, а не в конец листа."""
    text = "\n".join([line_1c("999", deadline="01.01.2026 09:00"),
                      line_1c("888", deadline="10.08.2026 12:30")])
    result = at.sync(env["xlsx"], text, say=quiet)

    assert sorted(result["new"]) == ["888", "999"]
    # 999 — раньше всех; 888 (10.08) — между 444 (05.07) и 222 (20.09).
    assert numbers(load(env)) == ["999", "444", "888", "222", "111", "333", "555\n556"]


def test_new_tender_fields_written(env):
    text = line_1c("999", customer="Новый заказчик", etp="roseltorg",
                   name="Вал редуктора", deadline="10.08.2026 12:30",
                   comment="")
    at.sync(env["xlsx"], text, say=quiet)

    ws = load(env)
    r = row_of(ws, "999")
    assert ws.cell(r, at.COL_CUSTOMER).value == "Новый заказчик"
    assert ws.cell(r, at.COL_NAME).value == "Вал редуктора"
    assert ws.cell(r, at.COL_ETP).value == "roseltorg"
    # «Ссылка» новым строкам не заполняется — её ведут руками.
    assert ws.cell(r, at.COL_LINK).value is None
    # Комментарий в выгрузке пуст — ячейку не трогаем.
    assert ws.cell(r, at.COL_COMMENT).value is None


def test_new_row_comment_from_1c(env):
    """Комментарий из выгрузки 1С попадает в столбец G новой строки."""
    text = line_1c("999", comment="вал 877*1шт.до 12.00")
    at.sync(env["xlsx"], text, say=quiet)

    ws = load(env)
    assert ws.cell(row_of(ws, "999"), at.COL_COMMENT).value == "вал 877*1шт.до 12.00"


def test_sync_fetches_b2b_link_for_new_1c_row(env, monkeypatch):
    """Новый формат 1С -> строка B2B -> ссылка в сохранённом XLSX."""
    def lookup(_session, number):
        return [{"id": number, "заголовок": "Тест", "url":
                 "https://www.b2b-center.ru/market/test/tender-%s/" % number}]

    monkeypatch.setattr(b2b_links, "lookup_tender", lookup)
    collected = []
    monkeypatch.setattr(
        b2b_details, "collect_b2b_details",
        lambda _path, numbers, **_kwargs: collected.extend(numbers) or {
            "processed": list(numbers), "txt_written": len(numbers)})
    text = line_1c("4621463", name="Дробильно-размольное&#x20;",
                   etp="b2b-center", comment="Запчасти к дробилке")

    at.sync(env["xlsx"], text, fetch_b2b_links=True, say=quiet)

    ws = load(env)
    cell = ws.cell(row_of(ws, "4621463"), at.COL_LINK)
    assert cell.value == "Открыть"
    assert cell.hyperlink.target.endswith("/tender-4621463/")
    assert collected == ["4621463"]


def test_new_row_default_status(env):
    at.sync(env["xlsx"], line_1c("999"), say=quiet)

    ws = load(env)
    assert ws.cell(row_of(ws, "999"), at.COL_STATUS).value == at.DEFAULT_STATUS


def test_new_row_action_mark(env):
    """Вместо подсветки — пометка «Добавлен ДД.ММ.ГГГГ» в столбце «Действие»."""
    at.sync(env["xlsx"], line_1c("999"), say=quiet)

    ws = load(env)
    today = datetime.date.today().strftime("%d.%m.%Y")
    assert ws.cell(row_of(ws, "999"), at.COL_ACTION).value == at.ACTION_PREFIX + today


def test_new_row_not_filled(env):
    """Фон новой строки белый: цвет даёт условное форматирование, а не заливка."""
    at.sync(env["xlsx"], line_1c("999"), say=quiet)

    ws = load(env)
    r = row_of(ws, "999")
    assert all(ws.cell(r, c).fill.patternType is None for c in range(1, at.NCOL + 1))


def test_new_row_styled_like_others(env):
    """Оформление берётся из строки-шаблона: текстовый номер и служебный столбец L."""
    at.sync(env["xlsx"], line_1c("999"), say=quiet)

    ws = load(env)
    r = row_of(ws, "999")
    assert ws.cell(r, at.COL_NUMBER).number_format == "@"
    assert isinstance(ws.cell(r, at.COL_NUMBER).value, str)
    assert ws.cell(r, at.COL_NUMBER).font.sz == 10
    assert ws.cell(r, at.COL_SORT).font.sz == 8
    assert ws.row_dimensions[r].height == at.ROW_HEIGHT


def test_deadline_written_as_serial(env):
    """Срок пишется числом с нужным форматом — иначе ISNUMBER в подсветке не сработает."""
    at.sync(env["xlsx"], line_1c("999", deadline="10.08.2026 12:30"), say=quiet)

    ws = load(env)
    cell = ws.cell(row_of(ws, "999"), at.COL_DEADLINE)
    assert isinstance(cell.value, (int, float))
    assert cell.number_format == at.DEADLINE_FORMAT


def test_sort_formula_matches_row(env):
    """В «Сорт.» у каждой строки формула с её собственным номером."""
    at.sync(env["xlsx"], line_1c("999"), say=quiet)

    ws = load(env)
    for r in range(2, at.last_data_row(ws) + 1):
        assert ws.cell(r, at.COL_SORT).value == at.SORT_FORMULA.format(row=r)


# ------------------------- сохранность существующих данных -------------------------

def test_existing_rows_survive_rewrite(env):
    """Ручные данные, стиль, гиперссылка и примечание старых строк не теряются."""
    at.sync(env["xlsx"], line_1c("999"), say=quiet)

    ws = load(env)
    r = row_of(ws, "111")
    assert ws.cell(r, at.COL_CUSTOMER).value == "Первый заказчик"
    assert ws.cell(r, at.COL_COMMENT).value == "ручной комментарий"
    assert ws.cell(r, at.COL_CONTACTS).value == "Иванов, 8-900"
    assert ws.cell(r, at.COL_LINK).hyperlink.target == "https://example.com/1"
    assert "важное примечание" in ws.cell(r, at.COL_COMMENT).comment.text


def test_broken_hyperlink_survives(env):
    """Ссылка, у которой адрес в location, а не в target, тоже переезжает."""
    at.sync(env["xlsx"], line_1c("999"), say=quiet)

    ws = load(env)
    cell = ws.cell(row_of(ws, "222"), at.COL_LINK)
    assert cell.hyperlink.location == "https://example.com/broken"
    assert cell.font.bold is True


def test_existing_action_not_touched(env):
    """Пометки пользователя в «Действие» программа не переписывает."""
    at.sync(env["xlsx"], line_1c("999"), say=quiet)

    ws = load(env)
    assert ws.cell(row_of(ws, "111"), at.COL_ACTION).value == "Прежний статус: Податься"


def test_second_run_adds_nothing(env):
    """Повторный запуск с тем же буфером не дублирует строки."""
    text = "\n".join([line_1c("111"), line_1c("999")])
    at.sync(env["xlsx"], text, say=quiet)
    before = numbers(load(env))

    result = at.sync(env["xlsx"], text, say=quiet)

    assert result["new"] == []
    assert numbers(load(env)) == before


def test_merged_numbers_not_duplicated(env):
    """В одной ячейке «Номер» бывает два номера — оба считаются уже существующими."""
    text = "\n".join([line_1c("555"), line_1c("556")])

    result = at.sync(env["xlsx"], text, say=quiet)

    assert result["new"] == []
    assert result["existed"] == 2


# ------------------------- обвязка листа -------------------------

def test_table_ref_expanded(env):
    at.sync(env["xlsx"], line_1c("999"), say=quiet)

    table = load(env).tables[at.TABLE]
    assert table.ref == "A1:L7"
    assert table.autoFilter.ref == "A1:L7"


def test_conditional_formatting_ranges_expanded(env):
    """Диапазоны УФ пересчитаны, и второй блок стал цельным (без дыр)."""
    at.sync(env["xlsx"], line_1c("999"), say=quiet)

    ranges = {str(fmt.sqref): len(fmt.rules) for fmt in load(env).conditional_formatting}
    assert ranges == {"F2:F7": 1, "A2:K7": 1}


def test_data_validation_survives(env):
    """Выпадающий список статусов возвращается в файл после сохранения."""
    at.sync(env["xlsx"], line_1c("999"), say=quiet)

    xml = sheet_xml(env)
    assert b"x14:dataValidation" in xml
    assert "Справочники!$A$2:$A$11".encode() in xml
    ET.fromstring(xml)        # файл, который Excel сочтёт повреждённым, тут упадёт


def test_read_ext_lst_strips_xr_uid(env):
    """xr:uid ссылается на пространство имён, которого в файле openpyxl нет."""
    ext = at.read_ext_lst(env["xlsx"], 1)

    assert b"x14:dataValidation" in ext
    assert b"xr:uid" not in ext


def test_dashboard_ranges_widened(env):
    """Формулы других листов дотягиваются до новой последней строки."""
    at.sync(env["xlsx"], line_1c("999"), say=quiet)

    wb = openpyxl.load_workbook(env["xlsx"])
    assert wb["Дашборд"]["B2"].value == '=COUNTIF(Тендеры!$I$2:$I$7,"Победа")'


def test_wrong_layout_refused(env):
    """Разъехавшиеся заголовки — отказ до записи, файл не тронут."""
    wb = openpyxl.load_workbook(env["xlsx"])
    wb[at.SHEET]["G1"] = "Не то"
    wb.save(env["xlsx"])
    before = open(env["xlsx"], "rb").read()

    with pytest.raises(at.SyncError):
        at.sync(env["xlsx"], line_1c("999"), say=quiet)
    assert open(env["xlsx"], "rb").read() == before


# ------------------------- папки и ссылки -------------------------

def test_folder_created_and_linked(env):
    at.sync(env["xlsx"], line_1c("999"), say=quiet)

    assert os.path.isdir(os.path.join(env["work"], "999", at.SUBFOLDER))
    ws = load(env)
    cell = ws.cell(row_of(ws, "999"), at.COL_FOLDER)
    assert cell.value == at.FOLDER_TEXT
    assert cell.hyperlink.target.endswith("999")


def test_existing_folder_with_suffix_reused(env):
    """Папка «999 описание» считается папкой тендера 999 — новая не создаётся."""
    existing = os.path.join(env["work"], "999 поставка насосов")
    os.makedirs(existing)
    at.sync(env["xlsx"], line_1c("999"), say=quiet)

    assert not os.path.exists(os.path.join(env["work"], "999"))
    assert os.path.isdir(os.path.join(existing, at.SUBFOLDER))


# ------------------------- защита данных -------------------------

def test_backup_created_and_no_tmp_left(env):
    at.sync(env["xlsx"], line_1c("999"), say=quiet)

    backups = list((env["tmp"] / "backups").iterdir())
    assert len(backups) == 1
    copy = openpyxl.load_workbook(backups[0])[at.SHEET]
    assert at.last_data_row(copy) == len(DATA_ROWS) + 1     # состояние ДО правки
    assert not os.path.exists(env["xlsx"] + ".tmp")


def test_dry_run_changes_nothing(env):
    before = open(env["xlsx"], "rb").read()

    result = at.sync(env["xlsx"], line_1c("999"), dry_run=True, say=quiet)

    assert result["new"] == ["999"]
    assert open(env["xlsx"], "rb").read() == before
    assert not os.path.exists(os.path.join(env["work"], "999"))
    assert not os.path.exists(at.HISTORY_PATH)


def test_history_written_for_real_run(env):
    at.sync(env["xlsx"], line_1c("999"), say=quiet)

    records = at.load_history(at.HISTORY_PATH)
    assert len(records) == 1
    assert records[0]["numbers"] == ["999"]


# ------------------------- ожидаемые ошибки -------------------------

def test_missing_file(tmp_path):
    with pytest.raises(at.SyncError):
        at.sync(str(tmp_path / "нет.xlsx"), line_1c("999"), say=quiet)


def test_empty_clipboard(env):
    with pytest.raises(at.SyncError):
        at.sync(env["xlsx"], "   \n\n", say=quiet)


def test_unrecognized_clipboard(env):
    with pytest.raises(at.SyncError):
        at.sync(env["xlsx"], "просто текст без табуляций", say=quiet)


def test_missing_sheet(env, monkeypatch):
    monkeypatch.setattr(at, "SHEET", "Такого листа нет")
    with pytest.raises(at.SyncError):
        at.sync(env["xlsx"], line_1c("999"), say=quiet)
