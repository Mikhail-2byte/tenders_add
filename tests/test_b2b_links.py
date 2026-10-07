# -*- coding: utf-8 -*-
"""Офлайн-тесты режима «Ссылки B2B» — сеть всегда подменяется."""

import os

import openpyxl
import pytest
import requests

from tenders import b2b_links as b
from tenders import core


def quiet(*_):
    pass


def test_is_b2b():
    assert b.is_b2b("b2b-center")
    assert b.is_b2b("b2b")
    assert b.is_b2b("B2B-Center")
    assert not b.is_b2b("Icetrade.by")
    assert not b.is_b2b("etp.acron.ru")
    assert not b.is_b2b(None)
    assert not b.is_b2b("")


def test_detect_block():
    assert b.detect_block("... servicepipe.tech ...")
    assert b.detect_block("<js-challenge-loader></js-challenge-loader>")
    assert b.detect_block("Forbidden. If you are not a bot ID: 42")
    assert b.detect_block("Превышен максимальный лимит скорости просмотра страниц")
    assert b.detect_block('{"trades": []}') is None


@pytest.mark.parametrize("relative, expected", [
    ("/market/some-slug/tender-4559367/#btid=2&sqh=temp&tsid=1",
     b.BASE + "/market/some-slug/tender-4559367/"),
    ("/app/market-next/some-slug/tender-4559367/#btid=2&sqh=temp&tsid=2",
     b.BASE + "/app/market-next/some-slug/tender-4559367/"),
])
def test_parse_results_exact_match_and_clean_url(relative, expected):
    payload = {"trades": [
        {"trade_id": 455936, "description": "Похожий номер", "url": "/wrong"},
        {"trade_id": 4559367, "description": "  Запрос   цен  ", "url": relative},
    ]}

    assert b.parse_results(payload, "4559367") == [{
        "id": "4559367", "заголовок": "Запрос цен", "url": expected,
    }]


def test_parse_results_ignores_fuzzy_matches():
    payload = {"trades": [
        {"trade_id": 462146, "url": "/market/a/tender-462146/"},
        {"trade_id": 46214630, "url": "/market/b/tender-46214630/"},
    ]}
    assert b.parse_results(payload, "4621463") == []


@pytest.mark.parametrize("payload", [None, [], {}, {"trades": None}, {"trades": ["bad"]}])
def test_parse_results_rejects_unknown_structure(payload):
    with pytest.raises(b.ProtocolChanged):
        b.parse_results(payload, "4621463")


def test_parse_results_rejects_match_without_url():
    with pytest.raises(b.ProtocolChanged):
        b.parse_results({"trades": [{"trade_id": 4621463}]}, "4621463")


class FakeResponse:
    def __init__(self, payload=None, status=200, content_type="application/json",
                 text="", json_error=None):
        self._payload = payload
        self._json_error = json_error
        self.status_code = status
        self.headers = {"Content-Type": content_type}
        self.text = text

    def raise_for_status(self):
        if self.status_code >= 400:
            raise requests.HTTPError("HTTP %d" % self.status_code)

    def json(self):
        if self._json_error:
            raise self._json_error
        return self._payload


class FakeSession:
    def __init__(self, response):
        self.response = response
        self.calls = []

    def get(self, url, **kwargs):
        self.calls.append((url, kwargs))
        return self.response


def test_lookup_uses_new_api_and_searches_archived(monkeypatch):
    monkeypatch.setattr(b, "ПАУЗА", 0)
    response = FakeResponse({"trades": [{
        "trade_id": 4559367,
        "description": "Архивная закупка",
        "url": "/market/archive/tender-4559367/#tracking",
    }]})
    session = FakeSession(response)

    result = b.lookup_tender(session, "4559367")

    assert result[0]["url"] == b.BASE + "/market/archive/tender-4559367/"
    assert session.calls == [(b.SEARCH_API, {
        "params": {"query": "4559367", "macro_trade_type": "buy", "tab": "all"},
        "timeout": 40,
    })]


@pytest.mark.parametrize("status", [403, 429])
def test_lookup_stops_on_http_limit(monkeypatch, status):
    monkeypatch.setattr(b, "ПАУЗА", 0)
    with pytest.raises(b.Antibot):
        b.lookup_tender(FakeSession(FakeResponse(status=status)), "4621463")


def test_lookup_stops_on_antibot_html(monkeypatch):
    monkeypatch.setattr(b, "ПАУЗА", 0)
    response = FakeResponse(content_type="text/html", text="If you are not a bot")
    with pytest.raises(b.Antibot):
        b.lookup_tender(FakeSession(response), "4621463")


def test_lookup_rejects_non_json_response(monkeypatch):
    monkeypatch.setattr(b, "ПАУЗА", 0)
    response = FakeResponse(content_type="text/html", text="обычная HTML-страница")
    with pytest.raises(b.ProtocolChanged):
        b.lookup_tender(FakeSession(response), "4621463")


def test_lookup_rejects_broken_json(monkeypatch):
    monkeypatch.setattr(b, "ПАУЗА", 0)
    response = FakeResponse(json_error=ValueError("bad json"))
    with pytest.raises(b.ProtocolChanged):
        b.lookup_tender(FakeSession(response), "4621463")


@pytest.fixture
def workbook(tmp_path):
    path = tmp_path / "заметки.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = core.SHEET
    ws.append(core.HEADERS)

    rows = [
        ("101", "b2b-center", None),
        ("102", "B2B-Center", "https://example.com/already"),
        ("103", "Фабрикант", None),
    ]
    for row, (number, etp, link) in enumerate(rows, start=2):
        ws.cell(row, core.COL_CUSTOMER, "Заказчик " + number)
        ws.cell(row, core.COL_NUMBER, number)
        ws.cell(row, core.COL_ETP, etp)
        if link:
            ws.cell(row, core.COL_LINK, "Открыть")
            ws.cell(row, core.COL_LINK).hyperlink = link
    wb.save(path)
    return path


def test_fill_writes_only_empty_b2b_link_and_creates_backup(workbook, monkeypatch):
    monkeypatch.setattr(b, "lookup_tender", lambda _session, number: [{
        "id": number,
        "заголовок": "Тест",
        "url": b.BASE + "/market/test/tender-%s/" % number,
    }])

    result = b.fill_b2b_links(str(workbook), say=quiet)

    ws = openpyxl.load_workbook(workbook)[core.SHEET]
    assert result["filled"] == 1
    assert ws.cell(2, core.COL_LINK).hyperlink.target.endswith("/tender-101/")
    assert ws.cell(3, core.COL_LINK).hyperlink.target == "https://example.com/already"
    assert ws.cell(4, core.COL_LINK).value is None
    assert result["stopped"] is None
    assert os.path.exists(result["backup"])
    backup_ws = openpyxl.load_workbook(result["backup"])[core.SHEET]
    assert backup_ws.cell(2, core.COL_LINK).value is None
    assert not os.path.exists(str(workbook) + ".tmp")


def test_fill_dry_run_does_not_change_file(workbook, monkeypatch):
    before = workbook.read_bytes()
    monkeypatch.setattr(b, "lookup_tender", lambda _session, number: [{
        "id": number, "заголовок": "Тест", "url": b.BASE + "/market/test/",
    }])

    result = b.fill_b2b_links(str(workbook), dry_run=True, say=quiet)

    assert result["filled"] == 1
    assert result["backup"] is None
    assert workbook.read_bytes() == before
    assert not (workbook.parent / "backups").exists()


def test_fill_can_be_limited_to_exact_numbers(workbook, monkeypatch):
    wb = openpyxl.load_workbook(workbook)
    ws = wb[core.SHEET]
    ws.cell(5, core.COL_CUSTOMER, "Заказчик 104")
    ws.cell(5, core.COL_NUMBER, "104")
    ws.cell(5, core.COL_ETP, "b2b-center")
    wb.save(workbook)
    looked_up = []

    def lookup(_session, number):
        looked_up.append(number)
        return [{"id": number, "заголовок": "Тест",
                 "url": b.BASE + "/market/test/tender-%s/" % number}]

    monkeypatch.setattr(b, "lookup_tender", lookup)
    result = b.fill_b2b_links(str(workbook), numbers=["104"], say=quiet)

    ws = openpyxl.load_workbook(workbook)[core.SHEET]
    assert looked_up == ["104"]
    assert result["filled"] == 1
    assert ws.cell(2, core.COL_LINK).value is None
    assert ws.cell(5, core.COL_LINK).hyperlink.target.endswith("/tender-104/")


def test_fill_records_network_error(workbook, monkeypatch):
    def fail(_session, _number):
        raise requests.ConnectionError("offline")

    monkeypatch.setattr(b, "lookup_tender", fail)
    result = b.fill_b2b_links(str(workbook), say=quiet)

    assert result["filled"] == 0
    assert result["errors"] == ["101"]
    assert result["stopped"] is None


def test_fill_stops_on_protocol_change_and_saves_collected_links(workbook, monkeypatch):
    wb = openpyxl.load_workbook(workbook)
    ws = wb[core.SHEET]
    ws.cell(5, core.COL_CUSTOMER, "Заказчик 104")
    ws.cell(5, core.COL_NUMBER, "104")
    ws.cell(5, core.COL_ETP, "b2b-center")
    wb.save(workbook)

    def lookup(_session, number):
        if number == "104":
            raise b.ProtocolChanged("формат изменился")
        return [{"id": number, "заголовок": "Тест",
                 "url": b.BASE + "/market/test/tender-%s/" % number}]

    monkeypatch.setattr(b, "lookup_tender", lookup)
    result = b.fill_b2b_links(str(workbook), say=quiet)

    ws = openpyxl.load_workbook(workbook)[core.SHEET]
    assert result["filled"] == 1
    assert result["stopped"] == "формат изменился"
    assert ws.cell(2, core.COL_LINK).hyperlink.target.endswith("/tender-101/")
    assert ws.cell(5, core.COL_LINK).value is None
