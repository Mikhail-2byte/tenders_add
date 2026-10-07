# -*- coding: utf-8 -*-
"""
Проставление ссылок B2B-Center в заметки.xlsx по номеру тендера.

Для строк, где столбец «ЭТП» указывает на B2B-Center, а «Ссылка» пуста, находит
процедуру публичным поиском площадки по номеру (столбец «Номер») и ставит в столбец
«Ссылка» кликабельную «Открыть» на карточку тендера.

Почему это безопасно: используется только публичный JSON-поиск площадки — один запрос
на номер, без авторизации и без обхода антибота. Массовый сбор и документация площадки
здесь НЕ трогаются (они под антибот-защитой и правами доступа).

ВАЖНО: заметки.xlsx должен быть ЗАКРЫТ (кроме --dry-run). Обработка идёт с паузой
ПАУЗА секунд между запросами — не уменьшать, регламент площадки не допускает ботов.

Запуск из консоли для отладки:
    py -3.12 -m tenders.b2b_links [--dry-run] [--file ПУТЬ]
"""

import argparse
import os
import re
import time
import warnings
from urllib.parse import urldefrag, urljoin

import openpyxl
import requests

# ------------------------- НАСТРОЙКИ B2B -------------------------
BASE = "https://www.b2b-center.ru"
SEARCH_API = BASE + "/site/api/v1/market-search/"
UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"
ПАУЗА = 1.5  # сек между запросами — не уменьшать (регламент площадки)

# Столбец «ЭТП» строки считается b2b-шным, если содержит одну из этих подстрок
# (в нижнем регистре). В файле встречаются значения «b2b» и «b2b-center».
B2B_ETP_MARKERS = ("b2b", "в2в")

LINK_TEXT = "Открыть"  # текст ссылки в столбце «Ссылка»
# ----------------------------------------------------------------


class SyncError(Exception):
    """Ожидаемая ошибка, текст которой показывается пользователю."""


class Antibot(SyncError):
    """Площадка отдала антибот-страницу — дальше запросы слать нельзя."""


class ProtocolChanged(SyncError):
    """Публичный поиск B2B вернул ответ неизвестного формата."""


def is_b2b(etp):
    """Строка относится к площадке B2B-Center?"""
    s = str(etp or "").lower()
    return any(m in s for m in B2B_ETP_MARKERS)


def detect_block(тело):
    """Распознаёт антибот-страницы (перенос из B2B-PoC).

    Все три состояния приходят так, что их легко принять за пустой результат,
    поэтому распознаём явно: JS-челлендж ServicePipe, блокировку по IP и
    превышение лимита скорости (все — с HTTP 200 и без нормального содержимого).
    """
    if "servicepipe.tech" in тело or "js-challenge-loader" in тело:
        return "Раздел закрыт JS-проверкой антибота (ServicePipe)."
    if "If you are not a bot" in тело:
        return "Площадка отдала антибот-страницу Forbidden."
    if "Превышен максимальный лимит скорости просмотра страниц" in тело:
        return ("Площадка включила лимит скорости просмотра страниц "
                "(регламент не допускает ботов).")
    return None


def parse_results(payload, number):
    """Точные совпадения номера из JSON-поиска B2B.

    Похожие результаты игнорируются. Tracking-фрагмент после ``#`` не нужен для
    открытия карточки и нестабилен, поэтому в Excel сохраняется чистый URL.
    """
    if not isinstance(payload, dict) or not isinstance(payload.get("trades"), list):
        raise ProtocolChanged("B2B-Center изменил формат ответа поиска.")

    wanted = str(number).strip()
    procedures = []
    for trade in payload["trades"]:
        if not isinstance(trade, dict):
            raise ProtocolChanged("B2B-Center изменил формат списка процедур.")
        if str(trade.get("trade_id", "")).strip() != wanted:
            continue
        relative_url = trade.get("url")
        if not isinstance(relative_url, str) or not relative_url.strip():
            raise ProtocolChanged("B2B-Center не вернул адрес найденной процедуры.")
        clean_url = urldefrag(urljoin(BASE + "/", relative_url.strip()))[0]
        title = " ".join(str(trade.get("description") or "").split())
        procedures.append({"id": wanted, "заголовок": title, "url": clean_url})
    return procedures


def lookup_tender(session, number):
    """Найти процедуру по номеру. Возвращает список совпадений (0/1/несколько).

    Antibot — если площадка отдала защитную страницу (тогда цикл надо остановить).
    """
    time.sleep(ПАУЗА)
    r = session.get(
        SEARCH_API,
        params={"query": number, "macro_trade_type": "buy", "tab": "all"},
        timeout=40,
    )
    if r.status_code in (403, 429):
        raise Antibot("Площадка временно ограничила публичный поиск (HTTP %d)."
                      % r.status_code)

    content_type = r.headers.get("Content-Type", "").lower()
    if "json" not in content_type:
        obstacle = detect_block(r.text)
        if obstacle:
            raise Antibot(obstacle)
        r.raise_for_status()
        raise ProtocolChanged("B2B-Center вместо JSON вернул ответ неизвестного формата.")

    r.raise_for_status()
    try:
        payload = r.json()
    except ValueError:
        raise ProtocolChanged("B2B-Center вернул повреждённый JSON.")
    return parse_results(payload, number)


def fill_b2b_links(xlsx_path, dry_run=False, say=print, numbers=None):
    """Проставить ссылки B2B в пустые «Ссылка» у строк с ЭТП = B2B.

    numbers — необязательный точный список номеров. Без него сохраняется прежний
    режим обработки всех B2B-строк без ссылки.

    Возвращает dict-итог; ожидаемые проблемы поднимает как SyncError.
    """
    # Ленивый импорт core чтобы избежать циклических зависимостей
    from . import core

    if not os.path.exists(xlsx_path):
        raise core.SyncError("Не найден файл " + xlsx_path)

    if not dry_run:
        try:
            with open(xlsx_path, "r+b"):
                pass
        except PermissionError:
            raise core.SyncError("Файл заметки.xlsx сейчас ОТКРЫТ в Excel.\n"
                                 "Закройте его и запустите снова.")

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(xlsx_path)
    if core.SHEET not in wb.sheetnames:
        raise core.SyncError("В файле нет листа «%s».\n"
                             "Похоже, выбран старый файл заметок." % core.SHEET)
    ws = wb[core.SHEET]
    core.check_layout(ws)
    sheet_index = wb.sheetnames.index(core.SHEET) + 1
    ext = core.read_ext_lst(xlsx_path, sheet_index)
    last = core.last_data_row(ws)

    wanted = None
    if numbers is not None:
        wanted = {str(number).strip() for number in numbers if str(number).strip()}

    # Цели: ЭТП = B2B, «Ссылка» пуста, номер есть.
    targets = []
    for r in range(2, last + 1):
        if not is_b2b(ws.cell(r, core.COL_ETP).value):
            continue
        if ws.cell(r, core.COL_LINK).value not in (None, ""):
            continue
        nums = core.numbers_in_cell(ws.cell(r, core.COL_NUMBER).value)
        if not nums:
            continue
        number = next((num for num in nums if wanted is None or num in wanted), None)
        if number:
            targets.append((r, number))

    say("Строк B2B без ссылки:", len(targets))
    if not targets:
        return {"dry_run": dry_run, "filled": 0, "not_found": [],
                "ambiguous": [], "errors": [], "backup": None,
                "stopped": None}

    session = requests.Session()
    session.headers["User-Agent"] = UA

    to_write = {}            # строка -> url
    not_found, ambiguous, errors = [], [], []
    stopped = None
    for r, number in targets:
        try:
            matches = lookup_tender(session, number)
        except (Antibot, ProtocolChanged) as ex:
            stopped = str(ex)
            break
        except requests.RequestException as ex:
            errors.append(number)
            say("  ! ошибка сети по", number, "—", ex)
            continue
        if not matches:
            not_found.append(number)
            say("  ?", number, "— не найдено")
        elif len(matches) > 1:
            ambiguous.append((number, [m["id"] for m in matches]))
            say("  ~", number, "— неоднозначно (%d совпадений): %s"
                % (len(matches), ", ".join(m["id"] for m in matches)))
        else:
            to_write[r] = matches[0]["url"]
            say("  +", number, "->", matches[0]["url"])

    if stopped:
        say("ОСТАНОВЛЕНО:", stopped)
        say("Обработать остаток можно позже — запросы прекращены, чтобы не "
            "нарушать регламент площадки.")

    if dry_run:
        say("[ПРЕДПРОСМОТР] Файл не изменяется.")
        say("Будет проставлено ссылок:", len(to_write))
        return {"dry_run": True, "filled": len(to_write), "not_found": not_found,
                "ambiguous": ambiguous, "errors": errors, "backup": None,
                "stopped": stopped}

    if not to_write:
        say("Проставлять нечего.")
        return {"dry_run": False, "filled": 0, "not_found": not_found,
                "ambiguous": ambiguous, "errors": errors, "backup": None,
                "stopped": stopped}

    for r, url in to_write.items():
        core.set_folder_link(ws.cell(r, core.COL_LINK), url)

    # Резервная копия оригинала; без неё файл не трогаем.
    try:
        backup = core.make_backup(xlsx_path)
    except OSError as ex:
        raise core.SyncError("Не удалось создать резервную копию: %s\n"
                             "Файл НЕ изменён." % ex)

    # Атомарное сохранение + возврат выпадающего списка статусов (write_ext_lst).
    # Строк не добавляли, поэтому границы таблицы и last не меняются.
    tmp_path = xlsx_path + ".tmp"
    try:
        wb.save(tmp_path)
        core.write_ext_lst(tmp_path, sheet_index, ext, last)
        os.replace(tmp_path, xlsx_path)
    except PermissionError:
        core.remove_quiet(tmp_path)
        raise core.SyncError("Не удалось сохранить: файл открыт в Excel. "
                             "Закройте его и повторите.")
    except Exception:
        core.remove_quiet(tmp_path)
        raise

    say("-" * 50)
    say("Проставлено ссылок:", len(to_write))
    say("Не найдено:", len(not_found), "| неоднозначно:", len(ambiguous),
        "| ошибок сети:", len(errors))
    say("Резервная копия:", backup)
    say("Готово. Файл сохранён:", xlsx_path)
    return {"dry_run": False, "filled": len(to_write), "not_found": not_found,
            "ambiguous": ambiguous, "errors": errors, "backup": backup,
            "stopped": stopped}


def main():
    # Ленивый импорт core чтобы избежать циклических зависимостей
    from . import core
    core.fix_console()
    ap = argparse.ArgumentParser(
        description="Проставить ссылки B2B-Center в заметки.xlsx по номеру тендера")
    ap.add_argument("--dry-run", action="store_true",
                    help="показать, что будет проставлено, без записи файла")
    ap.add_argument("--file", metavar="ПУТЬ",
                    help="путь к xlsx (по умолчанию — выбранный в окне программы)")
    args = ap.parse_args()
    try:
        fill_b2b_links(args.file or core.load_xlsx_path(), dry_run=args.dry_run)
    except core.SyncError as ex:
        print(ex)


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        import traceback
        traceback.print_exc()
        print("ОШИБКА:", e)
