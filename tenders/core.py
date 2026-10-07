# -*- coding: utf-8 -*-
"""
Ядро синхронизации тендеров из 1С с заметки.xlsx (лист "Тендеры").

Основной способ запуска — GUI: «Тендеры.bat» (tenders/app.py). Этот модуль
содержит всю логику (sync) и остаётся запускаемым из консоли для отладки:

    py -3.12 -m tenders.core [--dry-run] [--file ПУТЬ]

Как работает sync():
1. Получает текст буфера обмена (список тендеров из 1С, поля через табуляцию).
2. Сравнивает Номера с тем, что уже есть в Excel; НОВЫЕ тендеры получают статус
   "В работе", пометку "Добавлен ДД.ММ.ГГГГ" в столбце "Действие" и папку
   C:\\Работа\\В работе\\<Номер>\\С ЭТП с кликабельной ссылкой в столбце "Папка".
3. Пересортировывает весь лист: сначала по группе статуса (лист "Справочники"),
   затем по сроку подачи. Новые тендеры сами встают наверх — работа идёт сверху вниз.

Что НЕ трогается:
- значения и оформление существующих строк (Ссылка, Наименование, комментарии,
  Действие, Контакты, гиперссылки, примечания) сохраняются как есть. Меняется
  только положение строки на листе — по правилу сортировки.

Защита данных:
- перед сохранением оригинал копируется в подпапку backups (хранятся последние
  BACKUP_KEEP копий);
- запись атомарная: сначала во временный файл, затем замена оригинала — при
  сбое посреди записи оригинал остаётся цел.

Журналы и настройки (папка data в корне проекта):
- add_tenders.log — человекочитаемый итог каждого боевого запуска;
- history.jsonl — история запусков для статистики (по JSON-строке на запуск);
- config.json — путь к рабочему xlsx, выбранный в окне программы.

ВАЖНО: файл заметки.xlsx должен быть ЗАКРЫТ в момент запуска (кроме --dry-run).
"""

import argparse
import io
import json
import os
import re
import shutil
import sys
import tempfile
import datetime
import warnings
import zipfile
from copy import copy

import openpyxl
from openpyxl.comments import Comment
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

__version__ = "2.1.0"

# ------------------------- НАСТРОЙКИ -------------------------
WORK_DIR  = r"C:\Работа\В работе"
SHEET     = "Тендеры"        # рабочий лист
REF_SHEET = "Справочники"    # статусы и группы сортировки
TABLE     = "Тендеры_база"   # умная таблица на листе «Тендеры»
SUBFOLDER = "С ЭТП"

# Сколько последних резервных копий хранить в подпапке backups.
BACKUP_KEEP = 10

# Рабочие файлы программы (настройки, журнал, история) — в папке data корня проекта.
ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

TEMPLATE_PREFIX = "Запрос_"

DATA_DIR = os.path.join(ROOT_DIR, "data")
LOG_PATH = os.path.join(DATA_DIR, "add_tenders.log")
HISTORY_PATH = os.path.join(DATA_DIR, "history.jsonl")

# Путь к рабочему xlsx: выбирается в окне программы и запоминается в config.json;
# пока не выбран — заметки.xlsx на Рабочем столе текущего пользователя.
CONFIG_PATH = os.path.join(DATA_DIR, "config.json")
DEFAULT_XLSX = os.path.join(os.path.expanduser("~"), "Desktop", "заметки.xlsx")

# Возможные значения статуса в выгрузке 1С — "якорь":
# статус+1 = наименование, статус+2 = ЭТП, статус+3 = комментарий.
# К списку статусов Excel (лист «Справочники») отношения не имеет.
STATUSES_1C = {"в работе", "подался", "отказ", "проиграли",
               "победа", "интересный", "интересные", "на согласовании"}

# Форматы даты "Окончание подачи" в выгрузке 1С (пробуем по порядку).
DATE_FORMATS = ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%d.%m.%Y")

# Столбцы листа "Тендеры" (1=A ... 12=L)
COL_CUSTOMER, COL_NAME, COL_NUMBER, COL_LINK    = 1, 2, 3, 4
COL_ETP, COL_DEADLINE, COL_COMMENT, COL_ACTION  = 5, 6, 7, 8
COL_STATUS, COL_CONTACTS, COL_FOLDER, COL_SORT  = 9, 10, 11, 12
NCOL = 12

# Заголовки строки 1: макет фиксирован умной таблицей и сверяется перед записью.
HEADERS = ["Заказчик", "Наименование тендера", "Номер", "Ссылка", "ЭТП",
           "Окончание подачи", "Комментарий", "Действие", "Статус",
           "Контакты", "Папка", "Сорт."]

# Формула служебного столбца L "Сорт." — номер группы статуса из «Справочников».
SORT_FORMULA = ("=IFERROR(INDEX(Справочники!$E$2:$E$11,"
                "MATCH($I{row},Справочники!$C$2:$C$11,0)),99)")

# Лист «Справочники»: C — статусы в порядке сортировки, E — номер группы.
REF_STATUS_COL, REF_GROUP_COL = 3, 5
REF_ROW_FIRST, REF_ROW_LAST = 2, 11

DEFAULT_STATUS = "В работе"   # статус, который получают новые тендеры
UNKNOWN_GROUP  = 99           # группа для статуса, которого нет в справочнике
ROW_HEIGHT     = 30.0         # высота строки данных на листе «Тендеры»

# Формат "Окончание подачи": значения меньше 1 (пустые/нулевые) показываются пустыми.
DEADLINE_FORMAT = r'[<1]"";dd\.mm\.yyyy\ hh:mm'

FOLDER_TEXT   = "Открыть"     # текст ссылки в столбце "Папка"
ACTION_PREFIX = "Добавлен "   # пометка в столбце "Действие" у новых строк
# ------------------------------------------------------------


class SyncError(Exception):
    """Ожидаемая ошибка, текст которой показывается пользователю."""


def load_xlsx_path():
    """Путь к рабочему файлу: из config.json, иначе заметки.xlsx на Рабочем столе."""
    try:
        with open(CONFIG_PATH, encoding="utf-8") as f:
            path = json.load(f).get("xlsx_path")
    except (OSError, ValueError, AttributeError):
        path = None
    return path or DEFAULT_XLSX


def ensure_parent_dir(path):
    """Создаём папку для файла (data/), если её ещё нет."""
    parent = os.path.dirname(path)
    if parent:
        os.makedirs(parent, exist_ok=True)


def save_xlsx_path(path, say=print):
    """Запоминаем выбранный файл в config.json (папка data)."""
    try:
        ensure_parent_dir(CONFIG_PATH)
        with open(CONFIG_PATH, "w", encoding="utf-8") as f:
            json.dump({"xlsx_path": path}, f, ensure_ascii=False, indent=2)
    except OSError as ex:
        say("Не удалось сохранить настройки:", ex)


def fix_console():
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass


def read_clipboard():
    """Читаем текст из буфера обмена встроенным tkinter (без внешних пакетов)."""
    import tkinter
    root = tkinter.Tk()
    root.withdraw()
    try:
        data = root.clipboard_get()
    except Exception:
        data = ""
    finally:
        root.destroy()
    return data


def sanitize_folder_name(name):
    bad = '<>:"/\\|?*'
    for ch in bad:
        name = name.replace(ch, "_")
    return name.strip().rstrip(".")


def norm_number(value):
    """Номер как строка: неразрывные пробелы и переводы строк -> обычный пробел,
    повторы пробелов схлопываются, края обрезаются — чтобы номер из 1С совпадал
    с номером, когда-то вбитым в Excel руками."""
    if value in (None, ""):
        return None
    s = re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()
    return s or None


def numbers_in_cell(value):
    """Номера тендеров из ячейки "Номер": в одной ячейке их бывает несколько,
    записанных через перевод строки."""
    if value in (None, ""):
        return []
    return [n for n in (norm_number(p) for p in str(value).split("\n")) if n]


def list_work_folders():
    """Список подпапок в WORK_DIR (имена)."""
    try:
        return [d for d in os.listdir(WORK_DIR)
                if os.path.isdir(os.path.join(WORK_DIR, d))]
    except FileNotFoundError:
        return []


def find_existing_folder(num, folders):
    """Ищем папку тендера по номеру: точное имя или имя вида '<номер> описание'.
    Возвращаем полный путь или None."""
    target = sanitize_folder_name(num)
    for d in folders:                       # точное совпадение
        if d == target:
            return os.path.join(WORK_DIR, d)
    for d in folders:                       # '<номер> ...' или '<номер>_...'
        if d.startswith(target + " ") or d.startswith(target + "_"):
            return os.path.join(WORK_DIR, d)
    return None


def set_folder_link(cell, path):
    """Кликабельная ссылка «Открыть» на папку тендера — абсолютным путём, как есть.

    Оформление ячейки не трогаем: у существующих строк оно своё, у новых приходит
    из строки-шаблона. Относительные пути с %20, которые видно в файле, пишет сам
    Excel, сокращая наши абсолютные при очередном ручном сохранении.
    """
    cell.value = FOLDER_TEXT
    cell.hyperlink = path


def ensure_folder(num, work_folders, created_list, say=print):
    """Возвращает путь к папке тендера, создавая её (с подпапкой С ЭТП) при отсутствии.
    В created_list добавляет путь, только если подпапка была реально создана."""
    path = find_existing_folder(num, work_folders)
    if path is None:
        path = os.path.join(WORK_DIR, sanitize_folder_name(num))
    sub = os.path.join(path, SUBFOLDER)
    newly = not os.path.isdir(sub)
    try:
        os.makedirs(sub, exist_ok=True)
    except Exception as ex:
        say("  не удалось создать папку для", num, "->", ex)
        return None
    bn = os.path.basename(path)
    if bn not in work_folders:
        work_folders.append(bn)
    if newly:
        created_list.append(path)
    return path


def build_request_workbook():
    """Создаёт бланк запроса без зависимости от внешнего файла-шаблона."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.sheet_view.zoomScale = 70
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4

    widths = {
        "A": 9.14, "B": 29.43, "C": 23, "D": 20.43, "E": 26,
        "F": 14.71, "G": 14.57, "H": 13, "I": 29.57, "J": 9.14,
        "K": 13, "L": 17.86, "M": 26.71, "N": 17.14, "O": 17.57,
        "P": 13, "Q": 24.43, "R": 21.86, "S": 22.29,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    heights = {
        1: 21, 2: 29.25, 3: 21.75, 4: 21.75, 5: 89.25, 6: 21,
        7: 70.5, 8: 38.25, 18: 21, 19: 36.75, 20: 21, 21: 21,
        23: 20.25, 24: 19.5, 25: 15.75, 26: 15.75,
    }
    for row, height in heights.items():
        ws.row_dimensions[row].height = height

    thin = Side(style="thin", color="000000")
    cell_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    base_font = Font(name="Times New Roman", size=12)

    for row in ws.iter_rows(min_row=1, max_row=27, min_col=1, max_col=19):
        for cell in row:
            cell.font = base_font
            cell.alignment = Alignment(vertical="center")

    for row in range(1, 5):
        for column in range(1, 20):
            ws.cell(row, column).border = Border(
                left=thin if column == 1 else None,
                right=thin if column == 19 else None,
                top=thin,
                bottom=thin,
            )

    ws["A2"] = "Коммерческое предложение"
    ws["A2"].font = Font(name="Times New Roman", size=20)
    ws["A3"] = "Номер заявки : "
    ws["A3"].font = Font(name="Times New Roman", size=14, bold=True, italic=True)
    ws["A4"] = "Tо : "
    ws["A4"].font = Font(name="Times New Roman", size=12, bold=True, italic=True)

    headers = [
        "序号\n №", "产品名称\nНаименование", "产品材质 Материал изделия",
        "产品材料的模拟 Аналог материала изделия", "技术描述\nНомер чертежа",
        "硬度 Твёрдость, Shor", "硬度 Твёрдость, НВ", "硬度 Твёрдость, НRC",
        "技术要求 Технические требования", "单位\nЕд", "数量\nКол-во",
        "单价\nЦена ¥", "总金额\nСтоимость¥", "重量1件/公斤 Вес 1 шт/кг",
        "总重量       Общий вес, кг", "Пошлина",
        "生产厂家                                                          Завод-изготовитель",
        "生产期 Срок производства", "备注 Примечание",
    ]
    blue_fill = PatternFill("solid", fgColor="0070C0")
    red_fill = PatternFill("solid", fgColor="FF0000")
    for column, value in enumerate(headers, 1):
        cell = ws.cell(5, column, value)
        cell.font = Font(name="Times New Roman", size=12, bold=True, color="FFFFFF")
        cell.fill = blue_fill
        cell.alignment = center
        cell.border = cell_border

    sequence = {
        "A6": 1, "B6": 2, "C6": "=B6+1", "D6": "=C6+1",
        "E6": "=D6+1", "F6": "=E6+1", "G6": "=F6+1", "H6": "=G6+1",
        "I6": "=H6+1", "J6": "=I6+1", "K6": "=J6+1", "L6": "=K6+1",
        "M6": "=L6+1", "N6": "=M6+1", "O6": "=N6+1", "Q6": "=O6+1",
        "R6": "=Q6+1", "S6": "=R6+1",
    }
    for column in range(1, 20):
        cell = ws.cell(6, column)
        cell.fill = red_fill
        cell.font = Font(name="Times New Roman", size=14, color="FFFFFF")
        cell.alignment = center
        cell.border = cell_border
    for coordinate, value in sequence.items():
        ws[coordinate] = value

    for row in range(7, 9):
        for column in range(1, 20):
            ws.cell(row, column).alignment = center
            ws.cell(row, column).border = cell_border
    ws["A7"] = 1
    ws["I7"] = "По чертежу 根据图纸"
    ws["J7"] = "Штука"
    ws["M7"] = "=L7*K7"
    ws["O7"] = "=K7*N7"
    ws["S7"] = '=IFERROR(L7/N7,"")'

    ws.merge_cells("B8:L8")
    ws["B8"] = "Итого :"
    ws["B8"].alignment = Alignment(horizontal="right", vertical="center")
    ws["B8"].font = Font(name="Times New Roman", size=12, color="FF0000")
    ws["M8"] = "=SUM(M7:M7)"
    ws["O8"] = "=SUM(O7:O7)"
    ws["M8"].font = Font(name="Times New Roman", size=16, bold=True, color="FF0000")
    ws["O8"].font = Font(name="Times New Roman", size=16, bold=True, color="FF0000")

    for coordinate in ("L7", "M7", "L8", "M8"):
        ws[coordinate].number_format = '#,##0.00'
    for coordinate in ("N7", "O7", "O8", "S7"):
        ws[coordinate].number_format = '#,##0.00'

    notes = {
        "B17": "Примечание :",
        "B18": "Срок изготовления  :  …... дней  после получения аванс на заводе",
        "B19": "Условия поставки : FCA, Тяньзинь, Китай (ИНКОТЕРМС в редакции 2010 года)",
        "B20": "Срок действительна : ... дней",
        "B21": "Условия оплаты : …..% предоплата, ….% до отгрузки.",
        "E25": "00.00.2025",
    }
    for coordinate, value in notes.items():
        ws[coordinate] = value
    ws["B17"].font = Font(name="Times New Roman", size=12, color="FF0000")
    for coordinate in ("B18", "B19", "B20", "B21"):
        ws[coordinate].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    for merged_range in (
            "B18:R18", "B19:R19", "B20:E20", "B21:R21",
            "E24:R24", "B27:K27"):
        ws.merge_cells(merged_range)

    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    return wb


def copy_request_template(num, folder_path, say=print):
    """Создаёт бланк запроса, не заменяя существующий файл."""

    safe_num = sanitize_folder_name(num)
    target_name = f"{TEMPLATE_PREFIX}{safe_num}.xlsx"
    target_path = os.path.join(folder_path, target_name)

    if os.path.exists(target_path):
        return target_path

    wb = None
    temp_path = None
    try:
        wb = build_request_workbook()
        fd, temp_path = tempfile.mkstemp(
            prefix=f".{target_name}.", suffix=".tmp", dir=folder_path)
        os.close(fd)
        wb.save(temp_path)
        if os.path.exists(target_path):
            return target_path
        os.replace(temp_path, target_path)
        temp_path = None
        say(f"  создан файл запроса: {target_name}")
        return target_path
    except Exception as ex:
        say("  не удалось создать файл запроса для", num, "->", ex)
        return None
    finally:
        if wb is not None:
            wb.close()
        if temp_path is not None:
            try:
                os.remove(temp_path)
            except OSError:
                pass


def parse_deadline(raw):
    """Дата из 1С -> datetime. Не распозналось — возвращаем исходную строку."""
    for fmt in DATE_FORMATS:
        try:
            return datetime.datetime.strptime(raw, fmt)
        except ValueError:
            pass
    return raw


def parse_row(line):
    """Разбираем строку из 1С (поля через TAB). Возвращаем dict или None."""
    fields = [f.strip() for f in line.split("\t")]
    if len(fields) < 6:
        return None

    number = norm_number(fields[4]) or ""
    customer = fields[5]
    if not number or number.lower() in ("номер", "да", "нет") or number.lower() in STATUSES_1C:
        return None

    deadline_raw = fields[2] if len(fields) > 2 else ""

    name, etp, comment = "", "", ""
    status_idx = None
    for i, f in enumerate(fields):
        if f.lower() in STATUSES_1C:
            status_idx = i
            break
    if status_idx is not None:
        if status_idx + 1 < len(fields):
            name = fields[status_idx + 1]
        if status_idx + 2 < len(fields):
            etp = fields[status_idx + 2]
        if status_idx + 3 < len(fields):
            comment = fields[status_idx + 3]
    else:
        if len(fields) > 10:
            name = fields[10]
        if len(fields) > 11:
            etp = fields[11]
        if len(fields) > 12:
            comment = fields[12]

    # name — наименование тендера из 1С, едет в столбец B "Наименование тендера".
    return {"number": number, "customer": customer, "etp": etp,
            "deadline": parse_deadline(deadline_raw), "name": name,
            "comment": comment}


def check_layout(ws):
    """Сверяем заголовки: sync() стирает всю область данных, и запись при
    разъехавшихся столбцах испортила бы файл."""
    actual = [str(ws.cell(row=1, column=c).value or "").strip() for c in range(1, NCOL + 1)]
    if actual != HEADERS:
        raise SyncError("Не совпадают заголовки листа «%s» — файл не тронут.\n"
                        "Ожидается: %s\nВ файле:   %s"
                        % (SHEET, " | ".join(HEADERS), " | ".join(actual)))


def last_data_row(ws):
    """Последняя строка данных: по границе умной таблицы, иначе — сканом."""
    table = ws.tables.get(TABLE)
    if table is not None:
        m = re.search(r"(\d+)$", table.ref)
        if m:
            return int(m.group(1))
    last = 1
    for r in range(2, ws.max_row + 1):
        if (ws.cell(row=r, column=COL_CUSTOMER).value not in (None, "") or
                ws.cell(row=r, column=COL_NUMBER).value not in (None, "")):
            last = r
    return last


def load_groups(wb, say=print):
    """Порядок сортировки с листа «Справочники»: {статус в нижнем регистре: группа}."""
    groups = {}
    if REF_SHEET in wb.sheetnames:
        ws = wb[REF_SHEET]
        for r in range(REF_ROW_FIRST, REF_ROW_LAST + 1):
            status = ws.cell(row=r, column=REF_STATUS_COL).value
            group = ws.cell(row=r, column=REF_GROUP_COL).value
            if status and isinstance(group, (int, float)) and not isinstance(group, bool):
                groups[str(status).strip().lower()] = int(group)
    if not groups:
        say("Лист «%s» не найден или пуст — сортировка только по сроку подачи." % REF_SHEET)
    return groups


# Строки без срока и с нераспознанным сроком уезжают в конец своей группы.
NO_DEADLINE = 10 ** 9
EXCEL_EPOCH = datetime.datetime(1899, 12, 30)


def deadline_key(value):
    """Срок подачи -> число для сортировки."""
    if isinstance(value, datetime.datetime):
        return (value - EXCEL_EPOCH).total_seconds() / 86400.0
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    return NO_DEADLINE


def snapshot_cell(cell):
    """Полный снимок ячейки: значение + стиль + гиперссылка + комментарий."""
    cm = None
    if cell.comment is not None:
        cm = (cell.comment.text, cell.comment.author)
    return {
        "value": cell.value,
        "number_format": cell.number_format,
        "font": copy(cell.font),
        "fill": copy(cell.fill),
        "border": copy(cell.border),
        "alignment": copy(cell.alignment),
        "protection": copy(cell.protection),
        # Объект целиком, а не .target: у части ссылок адрес лежит в location.
        "hyperlink": copy(cell.hyperlink),
        "comment": cm,
    }


def apply_snapshot(cell, s):
    cell.value = s["value"]
    cell.number_format = s["number_format"]
    cell.font = s["font"]
    cell.fill = s["fill"]
    cell.border = s["border"]
    cell.alignment = s["alignment"]
    cell.protection = s["protection"]
    if s["hyperlink"] is not None:       # setter сам поправит ref на новую строку
        cell.hyperlink = s["hyperlink"]
    cell.comment = Comment(*s["comment"]) if s["comment"] else None


def apply_style_only(cell, s):
    """Применяем из снимка только оформление, не трогая значение/ссылку/комментарий."""
    cell.number_format = s["number_format"]
    cell.font = copy(s["font"])
    cell.fill = copy(s["fill"])
    cell.border = copy(s["border"])
    cell.alignment = copy(s["alignment"])
    cell.protection = copy(s["protection"])


def clear_cell(cell):
    cell.value = None
    cell.hyperlink = None
    cell.comment = None
    cell.style = "Normal"


def build_template(rows):
    """Оформление новой строки: по каждому столбцу берём первую строку данных
    БЕЗ заливки — иначе утащим ручную «зебру» и цвета условного форматирования."""
    template = []
    for c in range(1, NCOL + 1):
        pick = None
        for row in rows:
            fill = row["cells"][c - 1]["fill"]
            if fill is None or fill.patternType is None:
                pick = row["cells"][c - 1]
                break
        template.append(pick if pick is not None else rows[0]["cells"][c - 1])
    return template


# --- Выпадающий список статусов -------------------------------------------
# Excel хранит его в расширении x14 внутри <extLst> листа. openpyxl это
# расширение не понимает и при сохранении вырезает — возвращаем его вручную.

def read_ext_lst(path, sheet_index):
    """Достаём <extLst> листа из исходного xlsx (bytes; b"" — если его нет)."""
    try:
        with zipfile.ZipFile(path) as z:
            xml = z.read("xl/worksheets/sheet%d.xml" % sheet_index)
    except (OSError, KeyError, zipfile.BadZipFile):
        return b""
    m = re.search(rb"<extLst>.*?</extLst>(?=</worksheet>)", xml, re.S)
    if m is None:
        return b""
    # xr:uid ссылается на пространство имён, которого в файле openpyxl нет:
    # без удаления Excel посчитает файл повреждённым.
    return re.sub(rb' xr:uid="[^"]*"', b"", m.group())


def write_ext_lst(path, sheet_index, ext, last_row):
    """Вшиваем <extLst> обратно в сохранённый xlsx (архив пересобирается)."""
    if not ext:
        return
    # Диапазон выпадающего списка должен покрывать все строки таблицы с запасом.
    sqref = "<xm:sqref>I2:I%d</xm:sqref>" % max(1994, last_row + 200)
    ext = re.sub(rb"<xm:sqref>I2:I\d+</xm:sqref>", sqref.encode(), ext)
    target = "xl/worksheets/sheet%d.xml" % sheet_index
    with zipfile.ZipFile(path) as zin:
        parts = [(i.filename, zin.read(i.filename)) for i in zin.infolist()]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in parts:
            if name == target:
                data = data.replace(b"</worksheet>", ext + b"</worksheet>")
            zout.writestr(name, data)
    with open(path, "wb") as f:
        f.write(buf.getvalue())


# Жёстко зашитые в формулах других листов границы диапазона листа «Тендеры».
TENDER_RANGE_RE = re.compile(r"(Тендеры!\$[A-Z]{1,3}\$\d+:\$[A-Z]{1,3}\$)(\d+)")


def widen_formulas(wb, last, say=print):
    """Раздвигаем диапазоны Тендеры!...$749 в формулах «Дашборда» и листов-витрин:
    иначе после роста таблицы они молча перестанут учитывать новые строки."""
    changed = 0
    for name in wb.sheetnames:
        if name == SHEET:
            continue
        for row in wb[name].iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, ArrayFormula):
                    new = TENDER_RANGE_RE.sub(lambda m: m.group(1) + str(last), value.text)
                    if new != value.text:
                        cell.value = ArrayFormula(value.ref, new)
                        changed += 1
                elif isinstance(value, str) and value.startswith("="):
                    new = TENDER_RANGE_RE.sub(lambda m: m.group(1) + str(last), value)
                    if new != value:
                        cell.value = new
                        changed += 1
    if changed:
        say("Диапазоны формул на других листах расширены до строки %d (формул: %d)"
            % (last, changed))


def make_backup(path):
    """Копия xlsx в подпапку backups рядом с файлом; храним BACKUP_KEEP последних."""
    bdir = os.path.join(os.path.dirname(path) or ".", "backups")
    os.makedirs(bdir, exist_ok=True)
    base, ext = os.path.splitext(os.path.basename(path))
    stamp = datetime.datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    dst = os.path.join(bdir, base + "_" + stamp + ext)
    shutil.copy2(path, dst)
    old = sorted(f for f in os.listdir(bdir)
                 if f.startswith(base + "_") and f.endswith(ext))
    for f in old[:-BACKUP_KEEP]:
        try:
            os.remove(os.path.join(bdir, f))
        except OSError:
            pass
    return dst


def remove_quiet(path):
    try:
        os.remove(path)
    except OSError:
        pass


def append_log(lines, say=print):
    """Дописываем итог запуска в add_tenders.log."""
    try:
        ensure_parent_dir(LOG_PATH)
        with open(LOG_PATH, "a", encoding="utf-8") as f:
            f.write("\n".join(lines) + "\n\n")
    except OSError as ex:
        say("Не удалось записать лог:", ex)


# ------------------------- ИСТОРИЯ / СТАТИСТИКА -------------------------

def record_history(numbers, xlsx_path, say=print):
    """Дописываем запись об успешном запуске в history.jsonl (для статистики)."""
    rec = {"ts": datetime.datetime.now().isoformat(timespec="seconds"),
           "added": len(numbers), "numbers": list(numbers), "file": xlsx_path}
    try:
        ensure_parent_dir(HISTORY_PATH)
        with open(HISTORY_PATH, "a", encoding="utf-8") as f:
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")
    except OSError as ex:
        say("Не удалось записать историю:", ex)


def load_history(path=None):
    """Читаем history.jsonl; битые строки молча пропускаем."""
    if path is None:
        path = HISTORY_PATH
    records = []
    try:
        with open(path, encoding="utf-8") as f:
            for ln in f:
                ln = ln.strip()
                if not ln:
                    continue
                try:
                    rec = json.loads(ln)
                except ValueError:
                    continue
                if isinstance(rec, dict) and "ts" in rec:
                    records.append(rec)
    except OSError:
        pass
    return records


def counts_by_day(records):
    """Свод по дням: {date: {"count": сколько добавлено, "numbers": [...]}}."""
    days = {}
    for rec in records:
        try:
            day = datetime.datetime.fromisoformat(rec["ts"]).date()
        except (KeyError, TypeError, ValueError):
            continue
        d = days.setdefault(day, {"count": 0, "numbers": []})
        d["count"] += int(rec.get("added", 0) or 0)
        d["numbers"].extend(rec.get("numbers", []))
    return days


# ------------------------- ОСНОВНАЯ ЛОГИКА -------------------------

def sync(xlsx_path, text, dry_run=False, fetch_b2b_links=False, say=print):
    """Синхронизация: текст буфера 1С -> xlsx + папки.

    Возвращает dict с итогами; ожидаемые проблемы поднимает как SyncError
    (текст готов для показа пользователю). say(*args) — вывод хода работы.

    Если fetch_b2b_links=True, после сохранения автоматически проставит ссылки,
    соберёт карточки и скачает документы новых тендеров B2B-Center.
    """
    if not os.path.exists(xlsx_path):
        raise SyncError("Не найден файл " + xlsx_path)

    # В режиме предпросмотра ничего не пишем, поэтому открытый Excel не мешает.
    if not dry_run:
        try:
            with open(xlsx_path, "r+b"):
                pass
        except PermissionError:
            raise SyncError("Файл заметки.xlsx сейчас ОТКРЫТ в Excel.\n"
                            "Закройте его и запустите снова.")

    lines = [ln for ln in (text or "").replace("\r\n", "\n").split("\n") if ln.strip()]
    if not lines:
        raise SyncError("Буфер обмена пуст. Скопируйте список из 1С и повторите.")

    # Разбор буфера, сохраняя порядок 1С, без повторов внутри буфера
    entries, seen_clip = [], set()
    for ln in lines:
        e = parse_row(ln)
        if not e:
            continue
        if e["number"] in seen_clip:
            continue
        seen_clip.add(e["number"])
        entries.append(e)

    if not entries:
        raise SyncError("Не удалось распознать строки тендеров в буфере обмена.\n"
                        "Проверьте, что скопированы строки из 1С (поля через табуляцию).")

    # openpyxl предупреждает, что не понимает x14-валидацию: возвращаем её сами
    # (write_ext_lst), а предупреждение в окно программы пускать незачем.
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(xlsx_path)
    if SHEET not in wb.sheetnames:
        raise SyncError("В файле нет листа «%s».\n"
                        "Похоже, выбран старый файл заметок." % SHEET)
    ws = wb[SHEET]
    check_layout(ws)
    ext = read_ext_lst(xlsx_path, wb.sheetnames.index(SHEET) + 1)

    last = last_data_row(ws)
    groups = load_groups(wb, say)

    # Снимок существующих строк данных
    existing = []
    for r in range(2, last + 1):
        cells = [snapshot_cell(ws.cell(r, c)) for c in range(1, NCOL + 1)]
        existing.append({
            "numbers": numbers_in_cell(ws.cell(r, COL_NUMBER).value),
            "status": ws.cell(r, COL_STATUS).value,
            "deadline": ws.cell(r, COL_DEADLINE).value,
            "height": ws.row_dimensions[r].height,
            "cells": cells,
        })

    existing_numbers = set()
    for row in existing:
        existing_numbers.update(row["numbers"])

    new_entries = [e for e in entries if e["number"] not in existing_numbers]
    new_numbers = [e["number"] for e in new_entries]
    existed = len(entries) - len(new_entries)

    # Итоговый порядок строк: группа статуса, затем срок подачи.
    # Третий элемент ключа — исходный порядок: при равных группе и сроке
    # существующие строки остаются на местах, а новые встают следом за ними.
    default_group = groups.get(DEFAULT_STATUS.lower(), UNKNOWN_GROUP)
    final = []
    for i, row in enumerate(existing):
        status = str(row["status"] or "").strip().lower()
        final.append((groups.get(status, UNKNOWN_GROUP),
                      deadline_key(row["deadline"]), i, "existing", row))
    for i, e in enumerate(new_entries):
        final.append((default_group, deadline_key(e["deadline"]),
                      10 ** 6 + i, "new", e))
    final.sort(key=lambda item: item[:3])

    write_last = 1 + len(final)
    row_of = {}                 # номер тендера -> строка, куда он встанет
    for pos, item in enumerate(final):
        if item[3] == "new":
            row_of[item[4]["number"]] = pos + 2

    if dry_run:
        say("[ПРЕДПРОСМОТР] Файл и папки не изменяются.")
        say("Новых тендеров:", len(new_entries))
        for e in new_entries:
            say("   +", e["number"], "—", e["customer"],
                "(встанет строкой %d)" % row_of[e["number"]])
        say("Уже было в файле:", existed)
        folders = list_work_folders()
        to_create = []
        for e in entries:
            path = find_existing_folder(e["number"], folders)
            if path is None:
                to_create.append(os.path.join(
                    WORK_DIR, sanitize_folder_name(e["number"]), SUBFOLDER))
            elif not os.path.isdir(os.path.join(path, SUBFOLDER)):
                to_create.append(os.path.join(path, SUBFOLDER))
        if to_create:
            say("Будут созданы папки:")
            for p in to_create:
                say("   ", p)
        return {"dry_run": True, "new": new_numbers, "existed": existed,
                "folders_to_create": to_create, "created_folders": [],
                "linked": 0, "backup": None}

    # Чистим всю область данных, затем пишем заново в новом порядке
    for r in range(2, max(last, write_last) + 1):
        for c in range(1, NCOL + 1):
            clear_cell(ws.cell(r, c))

    template = build_template(existing) if existing else None
    clip_numbers = {e["number"] for e in entries}
    work_folders = list_work_folders()
    created_folders = []
    mark = ACTION_PREFIX + datetime.date.today().strftime("%d.%m.%Y")
    linked = 0
    r = 2
    for _, _, _, kind, src in final:
        if kind == "existing":
            for c in range(1, NCOL + 1):
                apply_snapshot(ws.cell(r, c), src["cells"][c - 1])
            if src["height"] is not None:
                ws.row_dimensions[r].height = src["height"]
            for num in src["numbers"]:
                # Тендер из текущего списка 1С — создаём папку, если её нет.
                # Старые строки (нет в списке) — только ссылка, если папка уже есть.
                if num in clip_numbers:
                    path = ensure_folder(num, work_folders, created_folders, say)
                else:
                    path = find_existing_folder(num, work_folders)
                if path:
                    set_folder_link(ws.cell(r, COL_FOLDER), path)
                    linked += 1
                    break
        else:
            e = src
            path = ensure_folder(e["number"], work_folders, created_folders, say)
            if path:
                copy_request_template(e["number"], path, say)
            if template is not None:
                for c in range(1, NCOL + 1):
                    apply_style_only(ws.cell(r, c), template[c - 1])
                    # Заливку не наследуем: фон белый, цвет даёт условное форматирование.
                    ws.cell(r, c).fill = PatternFill()
            ws.cell(r, COL_CUSTOMER, e["customer"])
            ws.cell(r, COL_NAME, e["name"])
            ws.cell(r, COL_NUMBER, e["number"])
            ws.cell(r, COL_ETP, e["etp"])
            if e["comment"]:
                ws.cell(r, COL_COMMENT, e["comment"])
            ws.cell(r, COL_STATUS, DEFAULT_STATUS)
            ws.cell(r, COL_ACTION, mark)
            if e["deadline"] not in (None, ""):
                dcell = ws.cell(r, COL_DEADLINE, e["deadline"])
                # Формат ТОЛЬКО после значения: иначе openpyxl затрёт его своим.
                dcell.number_format = DEADLINE_FORMAT
            ws.row_dimensions[r].height = ROW_HEIGHT
            if path:
                set_folder_link(ws.cell(r, COL_FOLDER), path)
                linked += 1
        # Служебный столбец «Сорт.» — вычисляемый, ручных данных в нём нет.
        ws.cell(r, COL_SORT, SORT_FORMULA.format(row=r))
        r += 1

    # Раздвигаем границы умной таблицы, автофильтра и сохранённой сортировки.
    table = ws.tables.get(TABLE)
    if table is not None:
        table.ref = "A1:%s%d" % (get_column_letter(NCOL), write_last)
        if table.autoFilter is not None:
            table.autoFilter.ref = table.ref
        if table.sortState is not None:
            table.sortState.ref = "A2:%s%d" % (get_column_letter(NCOL), write_last)

    # Диапазоны условного форматирования. Пересобираем список целиком: присваивание
    # sqref на месте ломает сохранение (ключ словаря правил хэшируется по нему).
    saved = [(str(fmt.sqref), list(fmt.rules)) for fmt in ws.conditional_formatting]
    ws.conditional_formatting = ConditionalFormattingList()
    for sqref, rules in saved:
        # Подсветка срока — по столбцу F, окраска строки по статусу — по A:K.
        wide = ("F2:F%d" % write_last if sqref.startswith("F2")
                else "A2:%s%d" % (get_column_letter(COL_FOLDER), write_last))
        for rule in rules:
            ws.conditional_formatting.add(wide, rule)

    widen_formulas(wb, write_last, say)

    unknown = sorted({str(row["status"]).strip() for row in existing
                      if row["status"]
                      and str(row["status"]).strip().lower() not in groups})
    if unknown:
        say("Статусы вне справочника (такие строки уходят вниз):", ", ".join(unknown))

    # Резервная копия оригинала; без неё файл не трогаем.
    try:
        backup = make_backup(xlsx_path)
    except OSError as ex:
        raise SyncError("Не удалось создать резервную копию: %s\n"
                        "Файл НЕ изменён. Проверьте место на диске/права и повторите." % ex)

    # Атомарное сохранение: пишем во временный файл, затем заменяем оригинал.
    tmp_path = xlsx_path + ".tmp"
    try:
        wb.save(tmp_path)
        # Возвращаем в файл выпадающий список статусов — до подмены оригинала.
        write_ext_lst(tmp_path, wb.sheetnames.index(SHEET) + 1, ext, write_last)
        os.replace(tmp_path, xlsx_path)
    except PermissionError:
        remove_quiet(tmp_path)
        raise SyncError("Не удалось сохранить: файл открыт в Excel. Закройте его и повторите.")
    except Exception:
        remove_quiet(tmp_path)
        raise

    say("-" * 50)
    say("Добавлено новых тендеров:", len(new_entries))
    say("Уже было в файле:", existed)
    if new_entries:
        say("Новые номера:", ", ".join(new_numbers))
    say("Ссылок на папки проставлено (столбец %s): %d"
        % (get_column_letter(COL_FOLDER), linked))
    if created_folders:
        say("Созданы папки:")
        for f in created_folders:
            say("   ", f)
    say("Резервная копия:", backup)
    say("Готово. Файл сохранён:", xlsx_path)

    # Автоматически проставить ссылки и собрать данные только новых B2B-тендеров.
    b2b_result = None
    b2b_details_result = None
    if fetch_b2b_links and new_entries:
        from . import b2b_links
        b2b_numbers = [entry["number"] for entry in new_entries
                       if b2b_links.is_b2b(entry["etp"])]
        if b2b_numbers:
            say("Поиск ссылок B2B-Center...")
            try:
                b2b_result = b2b_links.fill_b2b_links(
                    xlsx_path, dry_run=False, say=say, numbers=b2b_numbers)
                say("B2B: проставлено ссылок %d" % b2b_result.get("filled", 0))
                from . import b2b_details
                b2b_details_result = b2b_details.collect_b2b_details(
                    xlsx_path, b2b_numbers, dry_run=False, say=say)
            except b2b_links.Antibot as ex:
                say("B2B: остановлено антиботом — %s" % ex)
            except Exception as ex:
                # Excel, ссылка и папка уже сохранены. Ошибка закрытых данных не
                # должна откатывать основное добавление тендера.
                say("B2B: данные и документы не собраны — %s" % ex)

    log = ["[%s] %s" % (datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"), xlsx_path),
           "добавлено: %d, уже было: %d" % (len(new_entries), existed)]
    if new_entries:
        log.append("новые: " + ", ".join(new_numbers))
    if created_folders:
        log.append("созданы папки: " + "; ".join(created_folders))
    log.append("бэкап: " + backup)
    append_log(log, say)
    record_history(new_numbers, xlsx_path, say)

    return {"dry_run": False, "new": new_numbers, "existed": existed,
            "created_folders": created_folders, "linked": linked, "backup": backup,
            "b2b_links": b2b_result, "b2b_details": b2b_details_result}


def main():
    fix_console()

    ap = argparse.ArgumentParser(
        description="Добавление тендеров из буфера обмена (1С) в заметки.xlsx")
    ap.add_argument("--dry-run", action="store_true",
                    help="показать, что будет добавлено, без записи файла и создания папок")
    ap.add_argument("--fetch-b2b-links", action="store_true",
                    help="после добавления проставить ссылки B2B-Center для новых тендеров")
    ap.add_argument("--file", metavar="ПУТЬ",
                    help="путь к xlsx (по умолчанию — выбранный в окне программы)")
    args = ap.parse_args()

    try:
        sync(args.file or load_xlsx_path(), read_clipboard(),
             dry_run=args.dry_run, fetch_b2b_links=args.fetch_b2b_links)
    except SyncError as ex:
        print(ex)


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        import traceback
        traceback.print_exc()
        print("ОШИБКА:", e)
