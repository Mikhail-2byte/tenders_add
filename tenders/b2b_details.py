# -*- coding: utf-8 -*-
"""Карточка, позиции и документация B2B-Center через обычный браузер.

Вход выполняет пользователь в видимом Chrome/Edge. Программа хранит только
браузерный профиль, не принимает логин/пароль и не обходит защиту площадки.
"""
from __future__ import annotations

import datetime
import json
import os
import re
import shutil
import time
import warnings
import zipfile
from pathlib import Path, PurePosixPath
from urllib.parse import parse_qs, urlparse

import openpyxl
import requests

from . import b2b_links
from . import core


BASE = "https://www.b2b-center.ru"
AGGREGATE_API = BASE + "/market/openapi/v2/trade/get_trade_aggregate/"
INFO_FILENAME = "Информация B2B.txt"
PROFILE_DIR = os.path.join(
    os.environ.get("LOCALAPPDATA", os.path.expanduser("~")),
    "Rinako", "Tenders", "b2b-profile")
LOGIN_TIMEOUT = 300
PROCEDURE_PAUSE = 5
SCHEMA_VERSION = 1


class DetailsError(Exception):
    """Ожидаемая ошибка сбора одной процедуры."""


class AuthRequired(DetailsError):
    """Пользователь не выполнил вход в отведённое время."""


class ProtocolChanged(DetailsError):
    """Страница больше не соответствует ожидаемой структуре."""


class Antibot(DetailsError):
    """Площадка ограничила дальнейшую обработку."""


def _clean(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _hint_title(value):
    try:
        return _clean(value["name"]["hint"]["title"])
    except (KeyError, TypeError):
        return ""


def _field_value(fields, name, default=""):
    value = fields.get(name, default)
    if isinstance(value, dict) and "value" in value:
        return value["value"]
    return value


def _json_value(value):
    """Оставить из служебной структуры API только полезное значение."""
    if isinstance(value, list):
        return [_json_value(item) for item in value]
    if not isinstance(value, dict):
        return value
    if "address" in value and isinstance(value["address"], dict):
        return value["address"].get("address_string") or _json_value(value["address"])
    if "value" in value:
        return _json_value(value["value"])
    if "option" in value:
        return _hint_title(value["option"]) or _json_value(value["option"])
    ignored = {"settings", "sys_name", "name"}
    return {key: _json_value(item) for key, item in value.items()
            if key not in ignored}


def normalize_aggregate(payload, number, url):
    """Публичный JSON карточки -> стабильная часть нашего TXT."""
    if not isinstance(payload, dict):
        raise ProtocolChanged("B2B вернул неожиданный ответ карточки.")
    status = payload.get("status")
    aggregate = payload.get("trade_aggregate")
    trade = aggregate.get("trade") if isinstance(aggregate, dict) else None
    if not isinstance(status, dict) or status.get("code") != 1 or not isinstance(trade, dict):
        raise ProtocolChanged("B2B изменил структуру JSON карточки.")
    if str(trade.get("id")) != str(number):
        raise ProtocolChanged("B2B вернул карточку другого номера.")

    fields = trade.get("fields_values") or {}
    settings = trade.get("settings_values") or {}
    firm = trade.get("firm") or {}
    if not isinstance(fields, dict) or not isinstance(settings, dict) or not isinstance(firm, dict):
        raise ProtocolChanged("B2B вернул неполную структуру карточки.")

    delivery = _field_value(fields, "delivery_address", [])
    delivery_addresses = []
    if isinstance(delivery, list):
        for item in delivery:
            if isinstance(item, dict):
                address = item.get("address") or {}
                text = address.get("address_string") if isinstance(address, dict) else ""
                if text:
                    delivery_addresses.append(_clean(text))

    extra = {key: _json_value(value) for key, value in fields.items()}
    return {
        "procedure": {
            "number": str(number),
            "title": _clean(_field_value(fields, "subject")),
            "url": url,
            "type": _hint_title(trade.get("type") or {}),
            "status": (_hint_title(aggregate.get("trade_view_status") or {})
                       or _hint_title(trade.get("status") or {})),
            "published_at": trade.get("date_published") or "",
            "modified_at": trade.get("date_modified") or "",
            "deadline": _field_value(settings, "main_stage_date_end"),
            "positions_count": aggregate.get("positions_count"),
            "currency": _json_value(fields.get("currency", "")),
            "total_price": "",
            "category": "",
        },
        "organizer": {
            "name": _clean(firm.get("short_name")),
            "full_name": _clean(firm.get("full_name")),
            "inn": _clean(firm.get("inn")),
            "kpp": _clean(firm.get("kpp")),
            "ogrn": _clean(firm.get("ogrn")),
            "okpo": _clean(firm.get("okpo")),
            "legal_address": _clean(firm.get("jury_address")),
            "profile_url": firm.get("url") or "",
        },
        "terms": {
            "held_for": "",
            "payment": _clean(_field_value(fields, "payment_terms")),
            "delivery": _clean(_field_value(fields, "delivery_terms")),
            "delivery_address": delivery_addresses,
            "submission_method": "",
            "participant_comment": "",
        },
        "extra_fields": extra,
    }


def fetch_trade_aggregate(session, number):
    try:
        response = session.post(
            AGGREGATE_API, json={"trade_id": int(number)}, timeout=40)
    except requests.RequestException:
        raise
    if response.status_code in (403, 429):
        raise Antibot("B2B ограничил доступ к карточке (HTTP %d)." % response.status_code)
    obstacle = b2b_links.detect_block(response.text)
    if obstacle:
        raise Antibot(obstacle)
    response.raise_for_status()
    try:
        return normalize_aggregate(response.json(), number, "")
    except ValueError as ex:
        raise ProtocolChanged("B2B вернул повреждённый JSON карточки.") from ex


def sanitize_filename(name, fallback="Документы B2B.zip"):
    name = os.path.basename(str(name or "").replace("\\", "/"))
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name).strip(" .")
    if not name:
        name = fallback
    stem, ext = os.path.splitext(name)
    return (stem[:160] + ext[:20]).strip(" .") or fallback


def safe_extract_zip(archive_path, destination):
    """Распаковать ZIP атомарно, не позволяя именам выйти из destination."""
    archive_path = Path(archive_path)
    destination = Path(destination)
    try:
        archive = zipfile.ZipFile(archive_path)
    except (OSError, zipfile.BadZipFile) as ex:
        raise DetailsError("Повреждённый ZIP-пакет B2B: %s" % ex) from ex

    with archive:
        prepared = []
        for info in archive.infolist():
            raw = info.filename.replace("\\", "/")
            parts = PurePosixPath(raw).parts
            is_symlink = (info.external_attr >> 16) & 0o170000 == 0o120000
            if (PurePosixPath(raw).is_absolute() or ".." in parts
                    or (parts and ":" in parts[0]) or is_symlink):
                raise DetailsError("Небезопасный путь внутри ZIP: %s" % info.filename)
            if info.is_dir() or not parts:
                continue
            target = destination.joinpath(*parts)
            prepared.append((info, target))

        extracted = []
        for info, target in prepared:
            target.parent.mkdir(parents=True, exist_ok=True)
            tmp = target.with_name(target.name + ".part")
            try:
                with archive.open(info) as source, open(tmp, "wb") as output:
                    shutil.copyfileobj(source, output)
                os.replace(tmp, target)
            finally:
                try:
                    tmp.unlink()
                except FileNotFoundError:
                    pass
            extracted.append(str(target.relative_to(destination)))
        return extracted


def atomic_write_json(path, payload):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".tmp")
    try:
        with open(tmp, "w", encoding="utf-8", newline="\n") as output:
            json.dump(payload, output, ensure_ascii=False, indent=2)
            output.write("\n")
            output.flush()
            os.fsync(output.fileno())
        os.replace(tmp, path)
    finally:
        try:
            tmp.unlink()
        except FileNotFoundError:
            pass


def load_previous(path):
    try:
        with open(path, encoding="utf-8") as source:
            value = json.load(source)
        return value if isinstance(value, dict) else None
    except (OSError, ValueError):
        return None


def document_fingerprint(documents):
    values = []
    for doc in documents or []:
        if not isinstance(doc, dict):
            continue
        values.append((str(doc.get("checksum") or ""), str(doc.get("url") or ""),
                       str(doc.get("name") or "")))
    return json.dumps(sorted(values), ensure_ascii=False, separators=(",", ":"))


def _relative(root, path):
    return os.path.relpath(str(path), str(root)).replace("/", "\\")


def _absolute(root, relative):
    return Path(root).joinpath(*str(relative or "").replace("\\", "/").split("/"))


def _parse_contacts(raw):
    raw = _clean(raw)
    if not raw or "Контакты не добавлены" in raw:
        return []
    emails = sorted(set(re.findall(r"[\w.+-]+@[\w.-]+\.[A-Za-zА-Яа-я]{2,}", raw)))
    phones = sorted(set(re.findall(r"(?:\+7|8)[\d\s()\-]{8,}\d", raw)))
    return [{"raw": raw, "emails": emails, "phones": [_clean(x) for x in phones]}]


REQUIREMENT_FIELDS = {
    "Цены подаются", "Валюта заявки", "Объём поставки", "Альтернативные заявки",
    "Аналоги", "Подача регистрационной ставки", "Понижение цен участниками",
    "Отзыв заявки", "Документация", "Коммерческое предложение/оферта (обязательно)",
}


def build_record(number, url, aggregate, browser_data, previous=None):
    if not isinstance(browser_data, dict):
        raise ProtocolChanged("Браузер не вернул данные карточки.")
    page_number = re.sub(r"\D", "", str(browser_data.get("number") or ""))
    if page_number and page_number != str(number):
        raise ProtocolChanged("В браузере открылась карточка другого номера.")
    if not browser_data.get("title") and not aggregate.get("procedure", {}).get("title"):
        raise ProtocolChanged("На карточке B2B не найдено название процедуры.")

    sections = browser_data.get("sections") or {}
    procedure = dict(aggregate.get("procedure") or {})
    organizer = dict(aggregate.get("organizer") or {})
    terms = dict(aggregate.get("terms") or {})
    procedure.update({
        "number": str(number),
        "url": browser_data.get("url") or url,
        "title": browser_data.get("title") or procedure.get("title", ""),
        "category": browser_data.get("category") or procedure.get("category", ""),
        "status": procedure.get("status") or browser_data.get("status", ""),
        "total_price": sections.get("Общая сумма закупки", procedure.get("total_price", "")),
    })
    organizer["name"] = sections.get("Организатор", organizer.get("name", ""))
    terms.update({
        "held_for": sections.get("Процедура проводится", terms.get("held_for", "")),
        "payment": sections.get("Условия оплаты", terms.get("payment", "")),
        "delivery": sections.get("Условия поставки / оказания услуг", terms.get("delivery", "")),
        "delivery_address": sections.get(
            "Адрес поставки / оказания услуг", terms.get("delivery_address", [])),
        "submission_method": sections.get("Способ проведения", ""),
        "participant_comment": sections.get("Комментарий для участников", ""),
    })
    requirements = {key: value for key, value in sections.items()
                    if key in REQUIREMENT_FIELDS or "обязательно" in key.lower()}
    modeled = {
        "Общая сумма закупки", "Организатор", "Процедура проводится", "Условия оплаты",
        "Условия поставки / оказания услуг", "Адрес поставки / оказания услуг",
        "Способ проведения", "Комментарий для участников",
    } | set(requirements)
    extra = dict(aggregate.get("extra_fields") or {})
    extra.update({key: value for key, value in sections.items() if key not in modeled})

    errors = list(browser_data.get("errors") or [])
    return {
        "schema_version": SCHEMA_VERSION,
        "procedure": procedure,
        "organizer": organizer,
        "terms": terms,
        "requirements": requirements,
        "extra_fields": extra,
        "contacts": _parse_contacts(browser_data.get("contacts")),
        "positions": browser_data.get("positions") or [],
        "documents": browser_data.get("documents") or {
            "package": None, "listed": [], "obsolete": []},
        "meta": {
            "fetched_at": datetime.datetime.now().astimezone().isoformat(timespec="seconds"),
            "layout": browser_data.get("layout") or "unknown",
            "complete": not errors,
            "errors": errors,
        },
    }


def _browser_executable():
    candidates = [
        os.path.join(os.environ.get("ProgramFiles", ""),
                     "Google", "Chrome", "Application", "chrome.exe"),
        os.path.join(os.environ.get("ProgramFiles(x86)", ""),
                     "Google", "Chrome", "Application", "chrome.exe"),
        os.path.join(os.environ.get("ProgramFiles", ""),
                     "Microsoft", "Edge", "Application", "msedge.exe"),
        os.path.join(os.environ.get("ProgramFiles(x86)", ""),
                     "Microsoft", "Edge", "Application", "msedge.exe"),
    ]
    return next((path for path in candidates if path and os.path.isfile(path)), None)


NEW_CARD_JS = r"""
() => {
  const clean = value => (value || '').replace(/\s+/g, ' ').trim();
  const title = clean(document.title.split(/ — Тендер №/)[0]);
  const sections = {};
  for (const h of document.querySelectorAll('h5')) {
    const key = clean(h.textContent);
    const block = h.closest('[data-xid]') || h.parentElement;
    if (!key || !block) continue;
    const clone = block.cloneNode(true);
    clone.querySelectorAll('h5,svg').forEach(x => x.remove());
    const value = clean(clone.textContent);
    if (value && !sections[key]) sections[key] = value;
  }
  const docsRoot = document.querySelector('[data-xid="delivery-widget-organizer-docs"]');
  const documents = [...(docsRoot || document).querySelectorAll('a[href*="download.html"]')]
    .map(a => {
      const url = new URL(a.getAttribute('href'), location.href);
      return {name: clean(a.textContent), url: url.href,
              checksum: url.searchParams.get('checksum') || '',
              size: clean(a.parentElement && a.parentElement.textContent).replace(clean(a.textContent), '').trim()};
    });
  const contactHeading = [...document.querySelectorAll('h4')]
    .find(h => clean(h.textContent) === 'Контакты');
  const contactBlock = contactHeading && (contactHeading.parentElement?.parentElement || contactHeading.parentElement);
  const header = document.querySelector('.trade-header-title, h1');
  const tradeId = document.querySelector('[data-xid="trade-id"]');
  const stage = document.querySelector('[data-xid="trade-stages"]');
  return {title, number: clean(tradeId && tradeId.textContent), url: location.href,
          category: clean(header?.parentElement?.previousElementSibling?.textContent),
          status: clean(stage && stage.textContent), sections,
          contacts: clean(contactBlock && contactBlock.textContent), documents};
}
"""


LEGACY_CARD_JS = r"""
() => {
  const clean = value => (value || '').replace(/\s+/g, ' ').trim();
  const sections = {};
  for (const row of document.querySelectorAll('tr')) {
    const cells = [...row.children].filter(x => /^(TD|TH)$/.test(x.tagName));
    if (cells.length !== 2) continue;
    const key = clean(cells[0].textContent).replace(/:$/, '');
    const value = clean(cells[1].textContent);
    if (key && value && key.length < 200 && !sections[key]) sections[key] = value;
  }
  const documents = [...document.querySelectorAll('a[href*="download.html"]')]
    .map(a => { const url = new URL(a.getAttribute('href'), location.href);
      return {name: clean(a.textContent), url: url.href,
              checksum: url.searchParams.get('checksum') || '', size: ''}; });
  const title = clean(document.title.split(/ — Тендер №/)[0]);
  const match = location.href.match(/tender-(\d+)/) || location.href.match(/[?&]id=(\d+)/);
  const contacts = Object.entries(sections).filter(([k]) => /контакт/i.test(k))
    .map(([,v]) => v).join(' ');
  return {title, number: match ? match[1] : '', url: location.href,
          category: sections['Категория'] || '', status: sections['Статус'] || '',
          sections, contacts, documents};
}
"""


POSITIONS_JS = r"""
() => {
  const clean = value => (value || '').replace(/\s+/g, ' ').replace(/&nbsp;/g, '').trim();
  const root = document.querySelector('[data-xid="positions-table"]') || document;
  let headers = [...root.querySelectorAll('th')].map(x => clean(x.textContent)).filter(Boolean);
  if (!headers.length) {
    const first = [...root.querySelectorAll('tr')].find(r => r.querySelectorAll('th,td').length > 2);
    headers = first ? [...first.querySelectorAll('th,td')].map(x => clean(x.textContent)) : [];
  }
  const positions = [];
  let lot = '';
  for (const row of root.querySelectorAll('tbody tr, tr')) {
    const values = [...row.querySelectorAll(':scope > td')].map(x => clean(x.textContent));
    if (!values.length || values.every(x => !x)) continue;
    if (/^\d+$/.test(values[0]) && /Позиций в лоте/i.test(values[1] || '')) {
      lot = values[1].replace(/Позиций в лоте.*$/i, '').trim() || values[0];
      continue;
    }
    if (!/^\d+(?:\.\d+)?$/.test(values[0]) || /Показать ещё/i.test(values.join(' '))) continue;
    const raw = {};
    values.forEach((value, index) => { raw[headers[index] || `column_${index + 1}`] = value; });
    positions.push({lot, number: values[0], name: values[1] || '', raw});
  }
  return positions;
}
"""


class PlaywrightBrowser:
    def __init__(self, say=print):
        self.say = say
        self._playwright = None
        self.context = None
        self.page = None

    def __enter__(self):
        try:
            from playwright.sync_api import sync_playwright
        except ImportError as ex:
            raise DetailsError(
                "Не установлен Playwright. Выполните: py -3.12 -m pip install -r requirements.txt") from ex
        executable = _browser_executable()
        if not executable:
            raise DetailsError("Не найден установленный Google Chrome или Microsoft Edge.")
        Path(PROFILE_DIR).mkdir(parents=True, exist_ok=True)
        self._playwright = sync_playwright().start()
        try:
            self.context = self._playwright.chromium.launch_persistent_context(
                PROFILE_DIR, executable_path=executable, headless=False,
                accept_downloads=True, viewport=None, args=["--start-maximized"])
        except Exception:
            self._playwright.stop()
            self._playwright = None
            raise
        self.page = self.context.pages[0] if self.context.pages else self.context.new_page()
        return self

    def __exit__(self, *_):
        try:
            if self.context:
                self.context.close()
        finally:
            if self._playwright:
                self._playwright.stop()

    def _check_page(self, response=None):
        if response is not None and response.status in (403, 429):
            raise Antibot("B2B ограничил доступ (HTTP %d)." % response.status)
        text = self.page.locator("body").inner_text(timeout=15000)[:12000]
        obstacle = b2b_links.detect_block(text)
        if obstacle:
            raise Antibot(obstacle)

    def _ensure_login(self, url):
        def logged_in_page():
            # Форма входа B2B иногда завершает авторизацию в новой вкладке.
            # Следим за всем persistent context, а не только за исходной карточкой.
            for candidate in reversed(self.context.pages):
                try:
                    if (not candidate.is_closed()
                            and candidate.locator('[data-xid="header-logout"]').count()):
                        return candidate
                except Exception:
                    continue
            return None

        authenticated = logged_in_page()
        if authenticated is not None:
            self.page = authenticated
            if self.page.url != url:
                self.page.goto(url, wait_until="domcontentloaded", timeout=60000)
            return
        login = self.page.locator('a[href*="/auth/openid/authorize/"]').first
        if login.count() and login.is_visible():
            # Штатная ссылка сайта сохраняет redirect_uri обратно на карточку.
            # Пользователь вводит данные сам уже на открывшейся форме.
            login.click()
        self.say("B2B: войдите в открывшемся браузере. Ожидание — до 5 минут.")
        deadline = time.monotonic() + LOGIN_TIMEOUT
        while time.monotonic() < deadline:
            if all(candidate.is_closed() for candidate in self.context.pages):
                raise AuthRequired("Окно B2B закрыто до завершения входа.")
            authenticated = logged_in_page()
            if authenticated is not None:
                self.page = authenticated
                self.page.goto(url, wait_until="domcontentloaded", timeout=60000)
                return
            time.sleep(1)
        raise AuthRequired("Вход в B2B не выполнен за 5 минут.")

    def _download_package(self, number, tender_dir, documents, previous):
        etp_dir = Path(tender_dir) / core.SUBFOLDER
        etp_dir.mkdir(parents=True, exist_ok=True)
        fingerprint = document_fingerprint(documents)
        prior_docs = (previous or {}).get("documents") or {}
        prior_package = prior_docs.get("package") or {}
        prior_path = _absolute(tender_dir, prior_package.get("path"))
        same = fingerprint and fingerprint == prior_docs.get("fingerprint")
        if same and prior_path.is_file() and prior_path.stat().st_size:
            extracted = prior_package.get("extracted") or []
            missing = [item for item in extracted
                       if not _absolute(etp_dir, item).is_file()]
            if missing and zipfile.is_zipfile(prior_path):
                extracted = safe_extract_zip(prior_path, etp_dir)
                return ({**prior_package, "extracted": extracted, "status": "restored"},
                        len(extracted), 0, len(documents))
            return (prior_package, 0, 0, len(documents))

        button = self.page.locator('[data-xid="download-trade-docs-button"]').first
        if not button.count():
            button = self.page.get_by_role("button", name="Скачать все документы").first
        if not button.count() or not button.is_visible():
            if documents:
                return ({"status": "unavailable", "path": "", "extracted": []},
                        0, 0, 0)
            return (None, 0, 0, 0)

        with self.page.expect_download(timeout=120000) as pending:
            button.click()
        download = pending.value
        name = sanitize_filename(download.suggested_filename,
                                 "B2B-%s-документы.zip" % number)
        archive = etp_dir / name
        tmp = archive.with_name(archive.name + ".part")
        try:
            download.save_as(str(tmp))
            if not tmp.is_file() or not tmp.stat().st_size:
                raise DetailsError("B2B вернул пустой пакет документов.")
            os.replace(tmp, archive)
        finally:
            try:
                tmp.unlink()
            except FileNotFoundError:
                pass

        extracted = []
        archive_type = "file"
        if zipfile.is_zipfile(archive):
            archive_type = "zip"
            extracted = safe_extract_zip(archive, etp_dir)
        package = {
            "name": name,
            "path": _relative(tender_dir, archive),
            "type": archive_type,
            "size": archive.stat().st_size,
            "status": "downloaded",
            "extracted": extracted,
        }
        return package, len(extracted), 1, 0

    def _new_positions(self, card_url):
        positions_url = card_url.rstrip("/") + "/positions/"
        response = self.page.goto(positions_url, wait_until="domcontentloaded", timeout=60000)
        self._check_page(response)
        self.page.locator("table").first.wait_for(state="visible", timeout=60000)

        toggles = self.page.locator('[data-xid^="positions-row-expand-lot:"]')
        for index in range(toggles.count()):
            toggle = toggles.nth(index)
            if toggle.locator('svg g[id="plus"]').count():
                toggle.click()
                self.page.wait_for_timeout(350)

        clicks = 0
        while clicks < 500:
            buttons = self.page.locator(
                '[data-xid^="positions-load-more-loadmore-positions:lot:"]')
            if not buttons.count():
                break
            buttons.first.click()
            clicks += 1
            self.page.wait_for_timeout(450)
        if clicks >= 500:
            raise ProtocolChanged("Не удалось завершить загрузку всех позиций B2B.")
        return self.page.evaluate(POSITIONS_JS)

    def _legacy_positions(self, card_url):
        separator = "&" if "?" in card_url else "?"
        response = self.page.goto(card_url + separator + "action=positions",
                                  wait_until="domcontentloaded", timeout=60000)
        self._check_page(response)
        self.page.locator("table").first.wait_for(state="visible", timeout=60000)
        return self.page.evaluate(POSITIONS_JS)

    def collect(self, number, url, tender_dir, previous=None):
        response = self.page.goto(url, wait_until="domcontentloaded", timeout=60000)
        self._check_page(response)
        self._ensure_login(url)
        self._check_page()
        layout = "market-next" if "/app/market-next/" in self.page.url else "classic"
        if layout == "market-next":
            self.page.locator('[data-xid="trade-id"]').wait_for(
                state="visible", timeout=60000)
            card = self.page.evaluate(NEW_CARD_JS)
        else:
            self.page.locator("table").first.wait_for(state="visible", timeout=60000)
            card = self.page.evaluate(LEGACY_CARD_JS)

        errors = []
        documents = card.pop("documents", [])
        try:
            package, extracted, downloaded, skipped = self._download_package(
                number, tender_dir, documents, previous)
        except Exception as ex:
            package = {"status": "error", "path": "", "extracted": [],
                       "error": str(ex)}
            extracted = downloaded = skipped = 0
            errors.append("Документы: %s" % ex)

        extracted_names = {}
        if package:
            for relative in package.get("extracted") or []:
                extracted_names[Path(relative).name.casefold()] = relative
        listed = []
        for document in documents:
            item = dict(document)
            match = extracted_names.get(sanitize_filename(item.get("name"), "").casefold())
            if match:
                item["status"] = "downloaded"
                item["local_path"] = os.path.join(core.SUBFOLDER, match)
            elif package and package.get("status") in ("downloaded", "restored"):
                item["status"] = "in_package"
                item["local_path"] = package.get("path", "")
            else:
                item["status"] = "listed"
                item["local_path"] = ""
            listed.append(item)
        try:
            positions = (self._new_positions(url) if layout == "market-next"
                         else self._legacy_positions(url))
        except Exception as ex:
            positions = []
            errors.append("Позиции: %s" % ex)

        prior_listed = ((previous or {}).get("documents") or {}).get("listed") or []
        current_keys = {(item.get("checksum"), item.get("url")) for item in documents}
        obsolete = [item for item in prior_listed
                    if (item.get("checksum"), item.get("url")) not in current_keys]
        card.update({
            "layout": layout,
            "positions": positions,
            "errors": errors,
            "documents": {
                "package": package,
                "listed": listed,
                "obsolete": obsolete,
                "fingerprint": document_fingerprint(documents),
            },
            "files_extracted": extracted,
            "archives_downloaded": downloaded,
            "skipped_unchanged": skipped,
        })
        return card


def _cell_link(cell):
    link = cell.hyperlink
    if link is not None:
        return link.target or link.location or ""
    value = str(cell.value or "").strip()
    return value if value.startswith(("http://", "https://")) else ""


def _read_targets(xlsx_path, numbers):
    wanted = list(dict.fromkeys(str(number).strip() for number in numbers
                                if str(number).strip()))
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(xlsx_path)
    if core.SHEET not in wb.sheetnames:
        raise core.SyncError("В файле нет листа «%s»." % core.SHEET)
    ws = wb[core.SHEET]
    core.check_layout(ws)
    found_any, found_b2b, targets = set(), set(), {}
    for row in range(2, core.last_data_row(ws) + 1):
        row_numbers = core.numbers_in_cell(ws.cell(row, core.COL_NUMBER).value)
        matches = [number for number in wanted if number in row_numbers]
        if not matches:
            continue
        found_any.update(matches)
        if not b2b_links.is_b2b(ws.cell(row, core.COL_ETP).value):
            continue
        found_b2b.update(matches)
        for number in matches:
            targets[number] = {
                "number": number,
                "url": _cell_link(ws.cell(row, core.COL_LINK)),
                "row": row,
            }
    errors = []
    for number in wanted:
        if number not in found_any:
            errors.append({"number": number, "error": "номер не найден в Excel"})
        elif number not in found_b2b:
            errors.append({"number": number, "error": "у строки другая ЭТП"})
    return [targets[number] for number in wanted if number in targets], errors


def collect_b2b_details(xlsx_path, numbers, dry_run=False, say=print,
                        browser_factory=None):
    """Собрать карточки и файлы только для явно переданных B2B-номеров."""
    if not os.path.exists(xlsx_path):
        raise core.SyncError("Не найден файл " + xlsx_path)
    wanted = list(dict.fromkeys(re.findall(r"\d+", " ".join(map(str, numbers or [])))))
    result = {
        "dry_run": dry_run,
        "processed": [],
        "txt_written": 0,
        "archives_downloaded": 0,
        "files_extracted": 0,
        "skipped_unchanged": 0,
        "errors": [],
        "stopped": None,
        "auth_required": False,
    }
    if not wanted:
        return result

    targets, target_errors = _read_targets(xlsx_path, wanted)
    result["errors"].extend(target_errors)
    for item in target_errors:
        say("  !", item["number"], "—", item["error"])
    if dry_run:
        say("[ПРЕДПРОСМОТР] Браузер не открывается, файлы не изменяются.")
        say("Будет обработано B2B-процедур:", len(targets))
        return result

    missing_links = [target["number"] for target in targets if not target["url"]]
    if missing_links:
        say("B2B: ищем недостающие ссылки для", ", ".join(missing_links))
        b2b_links.fill_b2b_links(
            xlsx_path, dry_run=False, say=say, numbers=missing_links)
        targets, target_errors = _read_targets(xlsx_path, wanted)
        result["errors"].extend(item for item in target_errors
                                if item not in result["errors"])

    ready = []
    folders = core.list_work_folders()
    created = []
    for target in targets:
        if not target["url"]:
            result["errors"].append({
                "number": target["number"], "error": "ссылка B2B не найдена"})
            continue
        folder = core.ensure_folder(target["number"], folders, created, say)
        if not folder:
            result["errors"].append({
                "number": target["number"], "error": "не удалось создать папку"})
            continue
        target["folder"] = folder
        ready.append(target)
    if not ready:
        return result

    session = requests.Session()
    session.headers["User-Agent"] = b2b_links.UA
    factory = browser_factory or (lambda output: PlaywrightBrowser(output))
    try:
        with factory(say) as browser:
            for index, target in enumerate(ready):
                number = target["number"]
                say("B2B:", number, "— карточка, позиции и документы...")
                try:
                    info_path = Path(target["folder"]) / INFO_FILENAME
                    previous = load_previous(info_path)
                    aggregate = {
                        "procedure": {"number": number, "url": target["url"]},
                        "organizer": {}, "terms": {}, "extra_fields": {}}
                    aggregate_error = None
                    try:
                        aggregate = fetch_trade_aggregate(session, number)
                        aggregate["procedure"]["url"] = target["url"]
                    except Antibot:
                        raise
                    except Exception as ex:
                        aggregate_error = "JSON карточки: %s" % ex

                    browser_data = browser.collect(
                        number, target["url"], target["folder"], previous)
                    if aggregate_error:
                        browser_data.setdefault("errors", []).append(aggregate_error)
                    record = build_record(
                        number, target["url"], aggregate, browser_data, previous)
                    atomic_write_json(info_path, record)
                    result["processed"].append(number)
                    result["txt_written"] += 1
                    result["archives_downloaded"] += browser_data.get("archives_downloaded", 0)
                    result["files_extracted"] += browser_data.get("files_extracted", 0)
                    result["skipped_unchanged"] += browser_data.get("skipped_unchanged", 0)
                    say("  TXT:", info_path)
                except (AuthRequired, Antibot, ProtocolChanged):
                    raise
                except Exception as ex:
                    result["errors"].append({"number": number, "error": str(ex)})
                    say("  !", number, "—", ex)
                if index + 1 < len(ready):
                    time.sleep(PROCEDURE_PAUSE)
    except AuthRequired as ex:
        result["auth_required"] = True
        result["stopped"] = str(ex)
    except (Antibot, ProtocolChanged) as ex:
        result["stopped"] = str(ex)
    except DetailsError as ex:
        result["stopped"] = str(ex)
    except Exception as ex:
        result["errors"].append({"number": "", "error": str(ex)})

    if result["stopped"]:
        say("B2B остановлено:", result["stopped"])
    say("B2B: обработано %d, TXT %d, архивов %d, распаковано файлов %d"
        % (len(result["processed"]), result["txt_written"],
           result["archives_downloaded"], result["files_extracted"]))
    if result["errors"]:
        say("B2B: ошибок:", len(result["errors"]))
    return result
