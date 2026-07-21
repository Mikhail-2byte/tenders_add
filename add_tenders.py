# -*- coding: utf-8 -*-
"""
Ядро синхронизации тендеров из 1С с заметки.xlsx (лист "Лист1" = "В работе").

Основной способ запуска — GUI: «Тендеры.bat» (tender_app.py). Этот модуль
содержит всю логику (sync) и остаётся запускаемым из консоли для отладки:

    py -3.12 add_tenders.py [--dry-run] [--file ПУТЬ]

Как работает sync():
1. Получает текст буфера обмена (список тендеров из 1С, поля через табуляцию).
2. Сравнивает Номера с тем, что уже есть в Excel, и НОВЫЕ тендеры вставляет
   на нужное место (в порядке, как в 1С: сразу после предыдущего тендера из
   списка), заполняет колонки и создаёт папки C:\\Работа\\В работе\\<Номер>\\С ЭТП
   с кликабельной ссылкой в столбце "Папка".

Что НЕ трогается:
- уже существующие строки и все ручные данные (Ссылка, Наименование, комментарии,
  Действие, Контакты), их стили, форматы дат и гиперссылки сохраняются как есть.

Защита данных:
- перед сохранением оригинал копируется в подпапку backups (хранятся последние
  BACKUP_KEEP копий);
- запись атомарная: сначала во временный файл, затем замена оригинала — при
  сбое посреди записи оригинал остаётся цел.

Журналы и настройки (рядом со скриптом):
- add_tenders.log — человекочитаемый итог каждого боевого запуска;
- history.jsonl — история запусков для статистики (по JSON-строке на запуск);
- config.json — путь к рабочему xlsx, выбранный в окне программы.

ВАЖНО: файл заметки.xlsx должен быть ЗАКРЫТ в момент запуска (кроме --dry-run).
"""

import argparse
import json
import os
import shutil
import sys
import datetime
from copy import copy

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

__version__ = "1.0.0"

# ------------------------- НАСТРОЙКИ -------------------------
WORK_DIR  = r"C:\Работа\В работе"
SHEET     = "Лист1"
SUBFOLDER = "С ЭТП"

# Сколько последних резервных копий хранить в подпапке backups.
BACKUP_KEEP = 10

# Файл журнала запусков и файл истории для статистики (рядом со скриптом).
_HERE = os.path.dirname(os.path.abspath(__file__))
LOG_PATH = os.path.join(_HERE, "add_tenders.log")
HISTORY_PATH = os.path.join(_HERE, "history.jsonl")

# Путь к рабочему xlsx: выбирается в окне программы и запоминается в config.json;
# пока не выбран — заметки.xlsx на Рабочем столе текущего пользователя.
CONFIG_PATH = os.path.join(_HERE, "config.json")
DEFAULT_XLSX = os.path.join(os.path.expanduser("~"), "Desktop", "заметки.xlsx")

# Цвет подсветки тендеров, добавленных в текущем запуске (светло-зелёный).
# Подсветка снимается при следующем запуске, чтобы было видно только свежие.
HIGHLIGHT_RGB = "C6EFCE"

# Возможные значения статуса в выгрузке 1С — "якорь":
# поле сразу после статуса = ЭТП, следующее = описание.
STATUSES = {"в работе", "подался", "отказ", "проиграли",
            "победа", "интересный", "интересные", "на согласовании"}

# Форматы даты "Окончание подачи" в выгрузке 1С (пробуем по порядку).
DATE_FORMATS = ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%d.%m.%Y")

# Столбцы листа "Лист1" (1=A, 2=B, ...)
COL_CUSTOMER = 1   # A  Заказчик
COL_NUMBER   = 3   # C  Номер
COL_ETP      = 5   # E  ЭТП
COL_DEADLINE = 6   # F  Окончание подачи
COL_COMMENT  = 2   # B  Коментарий
FOLDER_HEADER = "Папка"
# ------------------------------------------------------------


class SyncError(Exception):
    """Ожидаемая ошибка, текст которой показывается пользователю."""


def load_xlsx_path():
    """Путь к рабочему файлу: из config.json, иначе заметки.xlsx на Рабочем столе."""
    try:
        with open(CONFIG_PATH, encoding="utf-8") as f:
            path = json.load(f).get("xlsx_path")
    except (OSError, ValueError, AttributeError):
        path = None
    return path or DEFAULT_XLSX


def save_xlsx_path(path, say=print):
    """Запоминаем выбранный файл в config.json рядом со скриптом."""
    try:
        with open(CONFIG_PATH, "w", encoding="utf-8") as f:
            json.dump({"xlsx_path": path}, f, ensure_ascii=False, indent=2)
    except OSError as ex:
        say("Не удалось сохранить настройки:", ex)


def fix_console():
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass


def read_clipboard():
    """Читаем текст из буфера обмена встроенным tkinter (без внешних пакетов)."""
    import tkinter
    root = tkinter.Tk()
    root.withdraw()
    try:
        data = root.clipboard_get()
    except Exception:
        data = ""
    finally:
        root.destroy()
    return data


def sanitize_folder_name(name):
    bad = '<>:"/\\|?*'
    for ch in bad:
        name = name.replace(ch, "_")
    return name.strip().rstrip(".")


def norm_number(value):
    """Номер как строка: неразрывные пробелы -> обычные, края обрезаются,
    чтобы номер из 1С совпадал с номером, когда-то вбитым в Excel руками."""
    if value in (None, ""):
        return None
    s = str(value).replace("\xa0", " ").strip()
    return s or None


def list_work_folders():
    """Список подпапок в WORK_DIR (имена)."""
    try:
        return [d for d in os.listdir(WORK_DIR)
                if os.path.isdir(os.path.join(WORK_DIR, d))]
    except FileNotFoundError:
        return []


def find_existing_folder(num, folders):
    """Ищем папку тендера по номеру: точное имя или имя вида '<номер> описание'.
    Возвращаем полный путь или None."""
    target = sanitize_folder_name(num)
    for d in folders:                       # точное совпадение
        if d == target:
            return os.path.join(WORK_DIR, d)
    for d in folders:                       # '<номер> ...' или '<номер>_...'
        if d.startswith(target + " ") or d.startswith(target + "_"):
            return os.path.join(WORK_DIR, d)
    return None


def make_highlight_fill():
    rgb = "FF" + HIGHLIGHT_RGB
    return PatternFill(start_color=rgb, end_color=rgb, fill_type="solid")


def is_highlight_fill(fill):
    """True, если ячейка залита нашим цветом подсветки (с прошлого запуска)."""
    try:
        if fill is None or fill.patternType != "solid":
            return False
        rgb = fill.fgColor.rgb
        return isinstance(rgb, str) and rgb.upper().endswith(HIGHLIGHT_RGB)
    except Exception:
        return False


def set_folder_link(cell, path):
    cell.value = "Открыть"
    cell.hyperlink = path
    cell.font = Font(color="0563C1", underline="single")


def ensure_folder(num, work_folders, created_list, say=print):
    """Возвращает путь к папке тендера, создавая её (с подпапкой С ЭТП) при отсутствии.
    В created_list добавляет путь, только если подпапка была реально создана."""
    path = find_existing_folder(num, work_folders)
    if path is None:
        path = os.path.join(WORK_DIR, sanitize_folder_name(num))
    sub = os.path.join(path, SUBFOLDER)
    newly = not os.path.isdir(sub)
    try:
        os.makedirs(sub, exist_ok=True)
    except Exception as ex:
        say("  не удалось создать папку для", num, "->", ex)
        return None
    bn = os.path.basename(path)
    if bn not in work_folders:
        work_folders.append(bn)
    if newly:
        created_list.append(path)
    return path


def parse_deadline(raw):
    """Дата из 1С -> (datetime, есть_ли_время). Не распозналось — (исходная строка, False)."""
    for fmt in DATE_FORMATS:
        try:
            return datetime.datetime.strptime(raw, fmt), "%H" in fmt
        except ValueError:
            pass
    return raw, False


def parse_row(line):
    """Разбираем строку из 1С (поля через TAB). Возвращаем dict или None."""
    fields = [f.strip() for f in line.split("\t")]
    if len(fields) < 6:
        return None

    number = norm_number(fields[4]) or ""
    customer = fields[5]
    if not number or number.lower() in ("номер", "да", "нет") or number.lower() in STATUSES:
        return None

    deadline_raw = fields[2] if len(fields) > 2 else ""

    etp, comment = "", ""
    status_idx = None
    for i, f in enumerate(fields):
        if f.lower() in STATUSES:
            status_idx = i
            break
    if status_idx is not None:
        if status_idx + 1 < len(fields):
            etp = fields[status_idx + 1]
        if status_idx + 2 < len(fields):
            comment = fields[status_idx + 2]
    else:
        if len(fields) > 10:
            etp = fields[10]
        if len(fields) > 11:
            comment = fields[11]

    deadline, has_time = parse_deadline(deadline_raw)

    return {"number": number, "customer": customer, "etp": etp,
            "deadline": deadline, "deadline_has_time": has_time,
            "comment": comment}


def ensure_folder_column(ws):
    """Находим столбец 'Папка'; если нет — создаём после последнего заголовка."""
    last_header_col = 0
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=1, column=c).value
        if v is not None and str(v).strip() != "":
            last_header_col = c
            if str(v).strip() == FOLDER_HEADER:
                return c
    col = last_header_col + 1
    cell = ws.cell(row=1, column=col, value=FOLDER_HEADER)
    cell.font = Font(bold=True)
    return col


def last_data_row(ws):
    last = 1
    for r in range(2, ws.max_row + 1):
        if (ws.cell(row=r, column=COL_CUSTOMER).value not in (None, "") or
                ws.cell(row=r, column=COL_NUMBER).value not in (None, "")):
            last = r
    return last


def snapshot_cell(cell):
    """Полный снимок ячейки: значение + стиль + гиперссылка + комментарий."""
    hl = cell.hyperlink.target if cell.hyperlink is not None else None
    cm = None
    if cell.comment is not None:
        cm = (cell.comment.text, cell.comment.author)
    return {
        "value": cell.value,
        "number_format": cell.number_format,
        "font": copy(cell.font),
        "fill": copy(cell.fill),
        "border": copy(cell.border),
        "alignment": copy(cell.alignment),
        "protection": copy(cell.protection),
        "hyperlink": hl,
        "comment": cm,
    }


def apply_snapshot(cell, s):
    cell.value = s["value"]
    cell.number_format = s["number_format"]
    cell.font = s["font"]
    cell.fill = s["fill"]
    cell.border = s["border"]
    cell.alignment = s["alignment"]
    cell.protection = s["protection"]
    cell.hyperlink = s["hyperlink"]      # строка-цель; None очищает
    cell.comment = Comment(*s["comment"]) if s["comment"] else None


def apply_style_only(cell, s):
    """Применяем из снимка только оформление, не трогая значение/ссылку/комментарий."""
    cell.number_format = s["number_format"]
    cell.font = copy(s["font"])
    cell.fill = copy(s["fill"])
    cell.border = copy(s["border"])
    cell.alignment = copy(s["alignment"])
    cell.protection = copy(s["protection"])


def clear_cell(cell):
    cell.value = None
    cell.hyperlink = None
    cell.comment = None
    cell.style = "Normal"


def make_backup(path):
    """Копия xlsx в подпапку backups рядом с файлом; храним BACKUP_KEEP последних."""
    bdir = os.path.join(os.path.dirname(path) or ".", "backups")
    os.makedirs(bdir, exist_ok=True)
    base, ext = os.path.splitext(os.path.basename(path))
    stamp = datetime.datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    dst = os.path.join(bdir, base + "_" + stamp + ext)
    shutil.copy2(path, dst)
    old = sorted(f for f in os.listdir(bdir)
                 if f.startswith(base + "_") and f.endswith(ext))
    for f in old[:-BACKUP_KEEP]:
        try:
            os.remove(os.path.join(bdir, f))
        except OSError:
            pass
    return dst


def remove_quiet(path):
    try:
        os.remove(path)
    except OSError:
        pass


def append_log(lines, say=print):
    """Дописываем итог запуска в add_tenders.log."""
    try:
        with open(LOG_PATH, "a", encoding="utf-8") as f:
            f.write("\n".join(lines) + "\n\n")
    except OSError as ex:
        say("Не удалось записать лог:", ex)


# ------------------------- ИСТОРИЯ / СТАТИСТИКА -------------------------

def record_history(numbers, xlsx_path, say=print):
    """Дописываем запись об успешном запуске в history.jsonl (для статистики)."""
    rec = {"ts": datetime.datetime.now().isoformat(timespec="seconds"),
           "added": len(numbers), "numbers": list(numbers), "file": xlsx_path}
    try:
        with open(HISTORY_PATH, "a", encoding="utf-8") as f:
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")
    except OSError as ex:
        say("Не удалось записать историю:", ex)


def load_history(path=None):
    """Читаем history.jsonl; битые строки молча пропускаем."""
    if path is None:
        path = HISTORY_PATH
    records = []
    try:
        with open(path, encoding="utf-8") as f:
            for ln in f:
                ln = ln.strip()
                if not ln:
                    continue
                try:
                    rec = json.loads(ln)
                except ValueError:
                    continue
                if isinstance(rec, dict) and "ts" in rec:
                    records.append(rec)
    except OSError:
        pass
    return records


def counts_by_day(records):
    """Свод по дням: {date: {"count": сколько добавлено, "numbers": [...]}}."""
    days = {}
    for rec in records:
        try:
            day = datetime.datetime.fromisoformat(rec["ts"]).date()
        except (KeyError, TypeError, ValueError):
            continue
        d = days.setdefault(day, {"count": 0, "numbers": []})
        d["count"] += int(rec.get("added", 0) or 0)
        d["numbers"].extend(rec.get("numbers", []))
    return days


# ------------------------- ОСНОВНАЯ ЛОГИКА -------------------------

def sync(xlsx_path, text, dry_run=False, say=print):
    """Синхронизация: текст буфера 1С -> xlsx + папки.

    Возвращает dict с итогами; ожидаемые проблемы поднимает как SyncError
    (текст готов для показа пользователю). say(*args) — вывод хода работы.
    """
    if not os.path.exists(xlsx_path):
        raise SyncError("Не найден файл " + xlsx_path)

    # В режиме предпросмотра ничего не пишем, поэтому открытый Excel не мешает.
    if not dry_run:
        try:
            with open(xlsx_path, "r+b"):
                pass
        except PermissionError:
            raise SyncError("Файл заметки.xlsx сейчас ОТКРЫТ в Excel.\n"
                            "Закройте его и запустите снова.")

    lines = [ln for ln in (text or "").replace("\r\n", "\n").split("\n") if ln.strip()]
    if not lines:
        raise SyncError("Буфер обмена пуст. Скопируйте список из 1С и повторите.")

    # Разбор буфера, сохраняя порядок 1С, без повторов внутри буфера
    entries, seen_clip = [], set()
    for ln in lines:
        e = parse_row(ln)
        if not e:
            continue
        if e["number"] in seen_clip:
            continue
        seen_clip.add(e["number"])
        entries.append(e)

    if not entries:
        raise SyncError("Не удалось распознать строки тендеров в буфере обмена.\n"
                        "Проверьте, что скопированы строки из 1С (поля через табуляцию).")

    wb = openpyxl.load_workbook(xlsx_path)
    if SHEET not in wb.sheetnames:
        raise SyncError("В файле нет листа " + SHEET)
    ws = wb[SHEET]

    folder_col = ensure_folder_column(ws)
    max_col = max(ws.max_column, folder_col)
    last = last_data_row(ws)

    # Снимок существующих строк данных
    existing = []
    for r in range(2, last + 1):
        cells = [snapshot_cell(ws.cell(r, c)) for c in range(1, max_col + 1)]
        existing.append({"number": norm_number(ws.cell(r, COL_NUMBER).value),
                         "cells": cells})

    # Снимаем подсветку с тендеров, добавленных в ПРОШЛЫЕ запуски: для каждой
    # колонки берём "обычную" заливку из первой неподсвеченной строки и
    # возвращаем её туда, где сейчас стоит зелёный цвет.
    normal_fill_by_col = {}
    for c in range(1, max_col + 1):
        for row in existing:
            f = row["cells"][c - 1]["fill"]
            if not is_highlight_fill(f):
                normal_fill_by_col[c] = copy(f)
                break
    for row in existing:
        for c in range(1, max_col + 1):
            s = row["cells"][c - 1]
            if is_highlight_fill(s["fill"]) and c in normal_fill_by_col:
                s["fill"] = copy(normal_fill_by_col[c])

    idx_by_num = {}
    for i, row in enumerate(existing):
        if row["number"] and row["number"] not in idx_by_num:
            idx_by_num[row["number"]] = i

    # Определяем новые тендеры и их "якорь" (после какой существующей строки вставить)
    inserts, new_entries = {}, []
    last_anchor = -1            # -1 = перед всеми строками
    for e in entries:
        if e["number"] in idx_by_num:
            last_anchor = idx_by_num[e["number"]]
        else:
            e["_anchor"] = last_anchor
            inserts.setdefault(last_anchor, []).append(e)
            new_entries.append(e)

    new_numbers = [e["number"] for e in new_entries]
    existed = len(entries) - len(new_entries)

    if dry_run:
        say("[ПРЕДПРОСМОТР] Файл и папки не изменяются.")
        say("Новых тендеров:", len(new_entries))
        for e in new_entries:
            say("   +", e["number"], "—", e["customer"])
        say("Уже было в файле:", existed)
        folders = list_work_folders()
        to_create = []
        for e in entries:
            path = find_existing_folder(e["number"], folders)
            if path is None:
                to_create.append(os.path.join(
                    WORK_DIR, sanitize_folder_name(e["number"]), SUBFOLDER))
            elif not os.path.isdir(os.path.join(path, SUBFOLDER)):
                to_create.append(os.path.join(path, SUBFOLDER))
        if to_create:
            say("Будут созданы папки:")
            for p in to_create:
                say("   ", p)
        return {"dry_run": True, "new": new_numbers, "existed": existed,
                "folders_to_create": to_create, "created_folders": [],
                "linked": 0, "backup": None}

    # Собираем итоговый порядок строк
    final = []                 # элементы: ('existing', row) или ('new', entry)
    for e in inserts.get(-1, []):
        final.append(("new", e))
    for i, row in enumerate(existing):
        final.append(("existing", row))
        for e in inserts.get(i, []):
            final.append(("new", e))

    write_last = 1 + len(final)

    # Чистим всю область данных, затем пишем заново
    for r in range(2, max(last, write_last) + 1):
        for c in range(1, max_col + 1):
            clear_cell(ws.cell(r, c))

    clip_numbers = {e["number"] for e in entries}
    work_folders = list_work_folders()
    created_folders = []
    linked = 0
    highlight_fill = make_highlight_fill()
    r = 2
    for kind, src in final:
        if kind == "existing":
            for c in range(1, max_col + 1):
                apply_snapshot(ws.cell(r, c), src["cells"][c - 1])
            num = src["number"]
            if num:
                # Тендер из текущего списка 1С — создаём папку, если её нет.
                # Старые строки (нет в списке) — только ссылка, если папка уже есть.
                if num in clip_numbers:
                    path = ensure_folder(num, work_folders, created_folders, say)
                else:
                    path = find_existing_folder(num, work_folders)
                if path:
                    set_folder_link(ws.cell(r, folder_col), path)
                    linked += 1
        else:
            e = src
            num = e["number"]
            path = ensure_folder(num, work_folders, created_folders, say)
            # Оформляем новую строку как соседнюю существующую (строку-«якорь»),
            # чтобы она не выделялась. Якорь -1 => берём первую строку данных.
            anchor = e.get("_anchor", -1)
            template = None
            if anchor >= 0:
                template = existing[anchor]
            elif existing:
                template = existing[0]
            if template is not None:
                for c in range(1, max_col + 1):
                    apply_style_only(ws.cell(r, c), template["cells"][c - 1])
            # Подсвечиваем новую строку светло-зелёным — видно, что добавлена сейчас.
            for c in range(1, folder_col + 1):
                ws.cell(r, c).fill = copy(highlight_fill)
            ws.cell(r, COL_CUSTOMER, e["customer"])
            ws.cell(r, COL_NUMBER, num)
            ws.cell(r, COL_ETP, e["etp"])
            dcell = ws.cell(r, COL_DEADLINE, e["deadline"])
            if isinstance(e["deadline"], datetime.datetime):
                dcell.number_format = ("DD.MM.YYYY HH:MM"
                                       if e.get("deadline_has_time") else "DD.MM.YYYY")
            ws.cell(r, COL_COMMENT, e["comment"])
            if path:
                set_folder_link(ws.cell(r, folder_col), path)
                linked += 1
        r += 1

    # Расширяем автофильтр на добавленные строки (если он задан на листе).
    if ws.auto_filter.ref:
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(folder_col), write_last)

    # Резервная копия оригинала; без неё файл не трогаем.
    try:
        backup = make_backup(xlsx_path)
    except OSError as ex:
        raise SyncError("Не удалось создать резервную копию: %s\n"
                        "Файл НЕ изменён. Проверьте место на диске/права и повторите." % ex)

    # Атомарное сохранение: пишем во временный файл, затем заменяем оригинал.
    tmp_path = xlsx_path + ".tmp"
    try:
        wb.save(tmp_path)
        os.replace(tmp_path, xlsx_path)
    except PermissionError:
        remove_quiet(tmp_path)
        raise SyncError("Не удалось сохранить: файл открыт в Excel. Закройте его и повторите.")
    except Exception:
        remove_quiet(tmp_path)
        raise

    say("-" * 50)
    say("Добавлено новых тендеров:", len(new_entries))
    say("Уже было в файле:", existed)
    if new_entries:
        say("Новые номера:", ", ".join(new_numbers))
    say("Ссылок на папки проставлено (столбец %s): %d"
        % (get_column_letter(folder_col), linked))
    if created_folders:
        say("Созданы папки:")
        for f in created_folders:
            say("   ", f)
    say("Резервная копия:", backup)
    say("Готово. Файл сохранён:", xlsx_path)

    log = ["[%s] %s" % (datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"), xlsx_path),
           "добавлено: %d, уже было: %d" % (len(new_entries), existed)]
    if new_entries:
        log.append("новые: " + ", ".join(new_numbers))
    if created_folders:
        log.append("созданы папки: " + "; ".join(created_folders))
    log.append("бэкап: " + backup)
    append_log(log, say)
    record_history(new_numbers, xlsx_path, say)

    return {"dry_run": False, "new": new_numbers, "existed": existed,
            "created_folders": created_folders, "linked": linked, "backup": backup}


def main():
    fix_console()

    ap = argparse.ArgumentParser(
        description="Добавление тендеров из буфера обмена (1С) в заметки.xlsx")
    ap.add_argument("--dry-run", action="store_true",
                    help="показать, что будет добавлено, без записи файла и создания папок")
    ap.add_argument("--file", metavar="ПУТЬ",
                    help="путь к xlsx (по умолчанию — выбранный в окне программы)")
    args = ap.parse_args()

    try:
        sync(args.file or load_xlsx_path(), read_clipboard(), dry_run=args.dry_run)
    except SyncError as ex:
        print(ex)


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        import traceback
        traceback.print_exc()
        print("ОШИБКА:", e)
